"""Page (references a view), Portal catalog/query, stamped report."""

from __future__ import annotations

import io

import pytest

from ._arnes_auth import usuario_logado
from .conftest import aguardar_execucao

pytestmark = pytest.mark.integration

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _view_pronta(cliente, h) -> str:
    """pipeline (city,amount) -> runs -> cube (measure total) -> published view. Returns view_id."""
    pid = (await cliente.post("/pipelines", json={"nome": "p"}, headers=h)).json()["id"]
    definicao = {
        "nos": [
            {
                "id": "saida",
                "tipo": "sql",
                "sql": "SELECT * FROM (VALUES ('SP',100),('SP',250),('RJ',90)) t(city, amount)",
            }
        ],
        "arestas": [],
    }
    await cliente.put(
        f"/pipelines/{pid}/rascunho", json={"definicao": definicao, "base_numero": 0}, headers=h
    )
    await cliente.post(
        f"/pipelines/{pid}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=h
    )
    eid = (await cliente.post(f"/pipelines/{pid}/executar", headers=h)).json()["id"]
    await aguardar_execucao(cliente, eid, h)  # snapshot ready before cube/view/page
    cid = (await cliente.post("/cubos", json={"nome": "C"}, headers=h)).json()["id"]
    cubodef = {
        "fonte": {"pipeline_id": pid, "no_id": "saida"},
        "dimensoes": [{"nome": "city", "coluna": "city"}],
        "medidas": [{"nome": "total", "expressao": "sum(amount)"}],
    }
    await cliente.put(
        f"/cubos/{cid}/rascunho", json={"definicao": cubodef, "base_numero": 0}, headers=h
    )
    await cliente.post(f"/cubos/{cid}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=h)
    vid = (await cliente.post("/views", json={"nome": "V"}, headers=h)).json()["id"]
    viewdef = {
        "cubo_id": cid,
        "linhas": ["city"],
        "valores": ["total"],
        "filtros_liberados": ["city"],
    }
    await cliente.put(
        f"/views/{vid}/rascunho", json={"definicao": viewdef, "base_numero": 0}, headers=h
    )
    await cliente.post(f"/views/{vid}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=h)
    return vid


async def _pagina(cliente, h, vid) -> str:
    pid = (await cliente.post("/paginas", json={"nome": "Sales"}, headers=h)).json()["id"]
    paginadef = {
        "blocos": [{"tipo": "pivot", "view_id": vid, "titulo": "By city"}],
        "filtros_pagina": [],
    }
    await cliente.put(
        f"/paginas/{pid}/rascunho", json={"definicao": paginadef, "base_numero": 0}, headers=h
    )
    await cliente.post(f"/paginas/{pid}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=h)
    return pid


async def test_pagina_catalogo_e_relatorio_carimbado(owner_engine, app_engine, cliente):
    _, he = await usuario_logado(
        owner_engine, cliente, email="espp@example.com", perfil="especialista"
    )
    vid = await _view_pronta(cliente, he)
    pag = await _pagina(cliente, he, vid)

    # Portal: the operator sees the catalog, queries and generates a report (read-only)
    _, ho = await usuario_logado(
        owner_engine, cliente, email="operp@example.com", perfil="operador"
    )
    cat = await cliente.get("/catalogo", headers=ho)
    assert cat.status_code == 200 and any(p["id"] == pag for p in cat.json())

    cons = await cliente.post(f"/paginas/{pag}/consultar", json={}, headers=ho)
    assert cons.status_code == 200
    bloco = cons.json()["blocos"][0]
    assert {x["city"]: x["total"] for x in bloco["resultado"]["linhas"]} == {"SP": 350, "RJ": 90}

    rel = await cliente.post(f"/paginas/{pag}/relatorio", json={}, headers=ho)
    assert rel.status_code == 200
    rid = rel.json()["id"]
    # trace chain: report -> ... -> execution
    assert rel.json()["trace"][0]["origem"]["execucao_id"]

    # the stamp is retrievable and immutable (same content)
    got = await cliente.get(f"/relatorios/{rid}", headers=ho)
    assert got.status_code == 200
    assert {x["city"]: x["total"] for x in got.json()["blocos"][0]["resultado"]["linhas"]} == {
        "SP": 350,
        "RJ": 90,
    }
    assert got.json()["trace"][0]["origem"]["execucao_id"]


async def test_listar_relatorios_e_exportar_xlsx(owner_engine, app_engine, cliente):
    _, he = await usuario_logado(
        owner_engine, cliente, email="esprel@example.com", perfil="especialista"
    )
    vid = await _view_pronta(cliente, he)
    pag = await _pagina(cliente, he, vid)

    _, ho = await usuario_logado(
        owner_engine, cliente, email="oprel@example.com", perfil="operador"
    )
    rel = await cliente.post(f"/paginas/{pag}/relatorio", json={}, headers=ho)
    assert rel.status_code == 200
    rid = rel.json()["id"]

    # listing: the newly created report shows up, with page name + criado_em (local time) + criado_por
    lista = await cliente.get("/relatorios", headers=ho)
    assert lista.status_code == 200
    item = next(r for r in lista.json() if r["id"] == rid)
    assert item["pagina"] == "Sales" and isinstance(item["versao"], int) and item["versao"] >= 1
    assert item["criado_em"] and "-03:00" in item["criado_em"]  # display time zone, explicit offset
    assert item["criado_por"]  # operator name

    # exportar.xlsx: 200 + xlsx content-type + body reopenable by openpyxl with the expected header
    xls = await cliente.get(f"/relatorios/{rid}/exportar.xlsx", headers=ho)
    assert xls.status_code == 200
    assert xls.headers["content-type"] == _XLSX_MIME
    assert 'filename="relatorio-' in xls.headers["content-disposition"]
    conteudo = xls.content
    assert conteudo  # non-empty body

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(conteudo))
    ws = wb["By city"]  # one sheet per block, name = sanitized block title
    cabecalho = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert cabecalho == ["city", "total"]
    dados = {r[0].value: r[1].value for r in ws.iter_rows(min_row=2)}
    assert dados == {"SP": 350, "RJ": 90}


async def test_exportar_pagina_xlsx_na_hora(owner_engine, app_engine, cliente):
    """POST /paginas/{id}/exportar.xlsx exports the query (with a released filter) without stamping."""
    _, he = await usuario_logado(
        owner_engine, cliente, email="esppx@example.com", perfil="especialista"
    )
    vid = await _view_pronta(cliente, he)
    pag = await _pagina(cliente, he, vid)

    _, ho = await usuario_logado(owner_engine, cliente, email="oppx@example.com", perfil="operador")
    xls = await cliente.post(
        f"/paginas/{pag}/exportar.xlsx",
        json={"filtros": [{"dimensao": "city", "op": "igual", "valor": "SP"}]},
        headers=ho,
    )
    assert xls.status_code == 200 and xls.headers["content-type"] == _XLSX_MIME
    assert 'filename="pagina-' in xls.headers["content-disposition"]

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xls.content))
    ws = wb["By city"]
    dados = {r[0].value: r[1].value for r in ws.iter_rows(min_row=2)}
    assert dados == {"SP": 350}  # filtro liberado aplicado


async def test_filtro_pagina_liberado_e_rbac(owner_engine, app_engine, cliente):
    _, he = await usuario_logado(
        owner_engine, cliente, email="espp2@example.com", perfil="especialista"
    )
    vid = await _view_pronta(cliente, he)
    pag = await _pagina(cliente, he, vid)

    _, ho = await usuario_logado(
        owner_engine, cliente, email="operp2@example.com", perfil="operador"
    )
    # filter on the released dimension
    r = await cliente.post(
        f"/paginas/{pag}/consultar",
        json={"filtros": [{"dimensao": "city", "op": "igual", "valor": "SP"}]},
        headers=ho,
    )
    assert {x["city"]: x["total"] for x in r.json()["blocos"][0]["resultado"]["linhas"]} == {
        "SP": 350
    }
    # the operator does NOT edit a page
    neg = await cliente.post("/paginas", json={"nome": "x"}, headers=ho)
    assert neg.status_code == 403


async def test_pagina_sem_bloco_422(owner_engine, app_engine, cliente):
    _, h = await usuario_logado(
        owner_engine, cliente, email="espp3@example.com", perfil="especialista"
    )
    pid = (await cliente.post("/paginas", json={"nome": "empty"}, headers=h)).json()["id"]
    r = await cliente.put(
        f"/paginas/{pid}/rascunho", json={"definicao": {"blocos": []}, "base_numero": 0}, headers=h
    )
    assert r.status_code == 422 and r.json()["erro"]["codigo"] == "definicao_invalida"


async def test_excluir_pagina_remove_e_audita(owner_engine, app_engine, cliente):
    from sqlalchemy import text

    _, h = await usuario_logado(
        owner_engine, cliente, email="delpag@example.com", perfil="especialista"
    )
    vid = await _view_pronta(cliente, h)
    pag = await _pagina(cliente, h, vid)  # creates a published version

    assert (await cliente.delete(f"/paginas/{pag}", headers=h)).status_code == 204
    assert (await cliente.get(f"/paginas/{pag}", headers=h)).status_code == 404
    assert pag not in [p["id"] for p in (await cliente.get("/catalogo", headers=h)).json()]
    async with owner_engine.connect() as conn:
        n_ver = (
            await conn.execute(
                text("SELECT count(*) FROM pagina_versoes WHERE pagina_id = :r"), {"r": pag}
            )
        ).scalar_one()
        n_perm = (
            await conn.execute(
                text(
                    "SELECT count(*) FROM permissoes WHERE recurso_tipo='pagina' AND recurso_id = :r"
                ),
                {"r": pag},
            )
        ).scalar_one()
        n_aud = (
            await conn.execute(
                text(
                    "SELECT count(*) FROM auditoria WHERE acao='excluir' AND recurso_tipo='pagina' AND recurso_id = :r"
                ),
                {"r": pag},
            )
        ).scalar_one()
    assert n_ver == 0 and n_perm == 0 and n_aud == 1


async def test_excluir_pagina_com_relatorio_carimbado_409(owner_engine, app_engine, cliente):
    # the stamp is immutable and part of the trace chain: a page with a report cannot be deleted
    _, h = await usuario_logado(
        owner_engine, cliente, email="delpag2@example.com", perfil="especialista"
    )
    vid = await _view_pronta(cliente, h)
    pag = await _pagina(cliente, h, vid)
    rel = await cliente.post(f"/paginas/{pag}/relatorio", json={}, headers=h)
    assert rel.status_code == 200

    r = await cliente.delete(f"/paginas/{pag}", headers=h)
    assert r.status_code == 409 and r.json()["erro"]["codigo"] == "pagina_com_relatorio"
    assert (await cliente.get(f"/paginas/{pag}", headers=h)).status_code == 200  # segue existindo


async def test_excluir_pagina_authz_negada(owner_engine, app_engine, cliente):
    _, dono = await usuario_logado(
        owner_engine, cliente, email="pdono@example.com", perfil="especialista"
    )
    pag = (await cliente.post("/paginas", json={"nome": "Protected"}, headers=dono)).json()["id"]
    _, outro = await usuario_logado(
        owner_engine, cliente, email="pintruso@example.com", perfil="especialista"
    )
    assert (await cliente.delete(f"/paginas/{pag}", headers=outro)).status_code == 403


async def test_operador_le_parametros_pela_pagina(owner_engine, app_engine, cliente):
    """The operator builds the report form from the PAGE, not from the pipeline.

    `GET /pipelines/{id}` is specialist/admin only: the operator used to get 403 and the Portal showed
    "Failed to open pipeline". `GET /paginas/{id}/parametros` returns only the released declarations.
    """
    _, he = await usuario_logado(
        owner_engine, cliente, email="esppar@example.com", perfil="especialista"
    )
    vid = await _view_pronta(cliente, he)

    pid = (await cliente.post("/pipelines", json={"nome": "with-params"}, headers=he)).json()["id"]
    definicao = {
        "nos": [{"id": "saida", "tipo": "sql", "sql": "SELECT :unit AS op"}],
        "arestas": [],
        "parametros": [
            {
                "nome": "unit",
                "tipo": "opcoes",
                "obrigatorio": True,
                "opcoes": ["E01", "E02"],
                "rotulo": "Unit (pipeline)",
                "valor_preview": "E01",
            },
            {"nome": "data_inicio", "tipo": "data"},
            {"nome": "not_released", "tipo": "texto"},
        ],
    }
    await cliente.put(
        f"/pipelines/{pid}/rascunho", json={"definicao": definicao, "base_numero": 0}, headers=he
    )
    await cliente.post(
        f"/pipelines/{pid}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=he
    )

    pag = (await cliente.post("/paginas", json={"nome": "With Params"}, headers=he)).json()["id"]
    paginadef = {
        "blocos": [{"tipo": "pivot", "view_id": vid, "titulo": "x"}],
        "filtros_pagina": [],
        "execucao": {
            "pipeline_id": pid,
            "operador": [
                {"tipo": "operacao", "param": "unit", "rotulo": "Unit"},
                {"tipo": "data", "param": "data_inicio"},
                {"tipo": "texto", "param": "gone_from_pipeline"},
            ],
        },
    }
    await cliente.put(
        f"/paginas/{pag}/rascunho", json={"definicao": paginadef, "base_numero": 0}, headers=he
    )
    await cliente.post(
        f"/paginas/{pag}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=he
    )

    _, ho = await usuario_logado(
        owner_engine, cliente, email="oppar@example.com", perfil="operador"
    )
    # without the page endpoint, the operator has no way to read the declaration (generic 403)
    assert (await cliente.get(f"/pipelines/{pid}", headers=ho)).status_code == 403

    r = await cliente.get(f"/paginas/{pag}/parametros", headers=ho)
    assert r.status_code == 200
    lista = r.json()
    # page order; a param missing from the pipeline is ignored; `not_released` does not leak
    assert [p["nome"] for p in lista] == ["unit", "data_inicio"]
    assert lista[0]["rotulo"] == "Unit"  # the page label wins over the pipeline one
    assert lista[0]["decl"]["tipo"] == "opcoes" and lista[0]["decl"]["obrigatorio"] is True
    assert lista[0]["decl"]["opcoes"] == ["E01", "E02"]
    assert "valor_preview" not in lista[0]["decl"]  # the preview value belongs to the Studio
    assert lista[1]["rotulo"] == "data_inicio"  # no label on the page -> the parameter name

    # generate a report -> the operator follows the execution they triggered themselves (Portal poll)
    exec_r = await cliente.post(
        f"/pipelines/{pid}/executar",
        json={"parametros": {"unit": "E01"}, "pagina_id": pag},
        headers=ho,
    )
    assert exec_r.status_code == 200
    eid = exec_r.json()["id"]
    assert (await cliente.get(f"/execucoes/{eid}", headers=ho)).status_code == 200

    # a page without an `execucao` block -> empty list (nothing to fill in), not an error
    sem = await _pagina(cliente, he, vid)
    vazio = await cliente.get(f"/paginas/{sem}/parametros", headers=ho)
    assert vazio.status_code == 200 and vazio.json() == []

    # a page that does not exist/is not published -> 404 in the standard envelope
    import uuid as _uuid

    nao = await cliente.get(f"/paginas/{_uuid.uuid4()}/parametros", headers=ho)
    assert nao.status_code == 404 and nao.json()["erro"]["codigo"] == "nao_encontrado"


async def test_execucao_do_studio_nao_vaza_pro_operador(owner_engine, app_engine, cliente):
    """An execution triggered by the Studio (no `pagina_id`, by another actor) stays closed to the operator."""
    _, he = await usuario_logado(
        owner_engine, cliente, email="espexec@example.com", perfil="especialista"
    )
    pid = (await cliente.post("/pipelines", json={"nome": "studio-only"}, headers=he)).json()["id"]
    definicao = {"nos": [{"id": "saida", "tipo": "sql", "sql": "SELECT 1 AS x"}], "arestas": []}
    await cliente.put(
        f"/pipelines/{pid}/rascunho", json={"definicao": definicao, "base_numero": 0}, headers=he
    )
    await cliente.post(
        f"/pipelines/{pid}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=he
    )
    eid = (await cliente.post(f"/pipelines/{pid}/executar", headers=he)).json()["id"]

    _, ho = await usuario_logado(
        owner_engine, cliente, email="opexec@example.com", perfil="operador"
    )
    neg = await cliente.get(f"/execucoes/{eid}", headers=ho)
    assert neg.status_code == 403 and neg.json()["erro"]["codigo"] == "acesso_negado"


async def test_catalogo_expoe_execucao(owner_engine, app_engine, cliente):
    """A published page WITH an `execucao` block returns it in /catalogo (the Portal runs it parameterized).
    The block round-trips through jsonb (draft -> publish); the pipeline_id may even be dangling."""
    _, he = await usuario_logado(
        owner_engine, cliente, email="espp4@example.com", perfil="especialista"
    )
    vid = await _view_pronta(cliente, he)
    fonte_pid = (await cliente.post("/pipelines", json={"nome": "source-exec"}, headers=he)).json()[
        "id"
    ]
    execucao = {
        "pipeline_id": fonte_pid,
        "operador": [
            {
                "tipo": "operacao",
                "param": "unit_id",
                "rotulo": "Unit",
                "mapa": "map-unit-system",
                "campo_id": "unit_id",
            },
            {
                "tipo": "mes",
                "rotulo": "Period month",
                "param_inicio": "data_inicio",
                "param_fim": "data_fim",
            },
        ],
        "derivacoes": [
            {
                "param": "sistema",
                "mapa": "map-unit-system",
                "campo_id": "unit_id",
                "campo_valor": "system_name",
                "normalizar": "sistema",
            },
            {
                "param": "projects",
                "mapa": "map-projects",
                "campo_id": "unit_id",
                "campo_valor": "project_names",
                "juntar": "|",
            },
        ],
    }
    pag = (await cliente.post("/paginas", json={"nome": "With Execution"}, headers=he)).json()["id"]
    paginadef = {
        "blocos": [{"tipo": "pivot", "view_id": vid, "titulo": "x"}],
        "filtros_pagina": [],
        "execucao": execucao,
    }
    await cliente.put(
        f"/paginas/{pag}/rascunho", json={"definicao": paginadef, "base_numero": 0}, headers=he
    )
    await cliente.post(
        f"/paginas/{pag}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=he
    )

    _, ho = await usuario_logado(
        owner_engine, cliente, email="operp4@example.com", perfil="operador"
    )
    cat = await cliente.get("/catalogo", headers=ho)
    assert cat.status_code == 200
    item = next(p for p in cat.json() if p["id"] == pag)
    assert item["execucao"] == execucao  # the published block round-trips intact
    assert item["execucao"]["pipeline_id"] == fonte_pid

    # a page without execucao: the key comes back null (not required)
    sem = await _pagina(cliente, he, vid)
    cat2 = await cliente.get("/catalogo", headers=ho)
    assert next(p for p in cat2.json() if p["id"] == sem)["execucao"] is None
