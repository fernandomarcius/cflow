"""GROUPED export through the API and the worker: a published grouped layout and three completed
executions yield ONE idempotent artifact with the three sheets, in order.

Also covers: request errors that name the sheet (409 for an unfinished execution, parts that do not
match, wrong layout type), publishing that validates the parts, RBAC at the edge, audit trail and
parts written in the SAME transaction, eager generation when the last base of the SAME run
completes (and only then), and failures that become state: the next attempt is a new record."""

from __future__ import annotations

import asyncio
import io
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from ._arnes_auth import usuario_logado
from .conftest import _WORKER, aguardar_execucao
from .test_relatorios_exportacoes import (
    _aguardar_exportacao,
    _arquivos_no_lake,
    _contar,
    _executar,
    _layout_publicado,
    _pipeline_publicado,
)

pytestmark = pytest.mark.integration

ABAS = ("SALES", "INVOICES", "PAYMENTS")
_SQL = {
    "SALES": "SELECT * FROM (VALUES ('S1', 10.5), ('S2', NULL)) t(sale, amount)",
    "INVOICES": "SELECT * FROM (VALUES ('I1', DATE '2026-08-01')) t(invoice, due_date)",
    "PAYMENTS": "SELECT * FROM (VALUES ('P1', 1), ('P2', 2), ('P3', 3)) t(payment, n)",
}
_COLUNAS = {
    "SALES": [{"coluna": "sale", "rotulo": "Sale"}, {"coluna": "amount", "formato": "moeda"}],
    "INVOICES": [{"coluna": "invoice"}, {"coluna": "due_date", "formato": "data"}],
    "PAYMENTS": [{"coluna": "payment"}, {"coluna": "n", "formato": "inteiro"}],
}


def _layout_base(pid: str, aba: str, padrao: bool = False) -> dict:
    return {
        "fonte": {"pipeline": pid, "no": "saida"},
        "aba": aba,
        "colunas": _COLUNAS[aba],
        "cabecalho_fixo": False,
        "autofiltro": False,
        "carimbo": False,
        "padrao": padrao,
    }


def _agrupado(layouts: dict[str, str], padrao: bool = False) -> dict:
    return {
        "tipo": "agrupado",
        "nome": "Bases agrupadas",
        "abas": [{"aba": aba, "layout": layouts[aba]} for aba in ABAS],
        "carimbo": False,
        "padrao": padrao,
    }


async def _tres_bases(cliente, h, sufixo: str) -> tuple[dict[str, str], dict[str, str]]:
    """Three pipelines with one `saida` node, plus a single-source layout for each. Returns (pids, layouts)."""
    pids, layouts = {}, {}
    for aba in ABAS:
        pids[aba] = await _pipeline_publicado(
            cliente, h, f"base-{aba}-{sufixo}", [{"id": "saida", "tipo": "sql", "sql": _SQL[aba]}]
        )
        layouts[aba] = await _layout_publicado(
            cliente, h, _layout_base(pids[aba], aba), nome=f"Base {aba} {sufixo}"
        )
    return pids, layouts


async def _pedir(cliente, h, lid: str, partes: dict[str, str]):
    return await cliente.post(
        "/relatorios/exportacoes", json={"layout_id": lid, "partes": partes}, headers=h
    )


# --- end-to-end flow --------------------------------------------------------------------------


async def test_grouped_is_idempotent_with_three_sheets_in_order(owner_engine, app_engine, cliente):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.agr@example.com", perfil="especialista"
    )
    _, ope = await usuario_logado(
        owner_engine, cliente, email="ope.agr@example.com", perfil="operador"
    )
    pids, layouts = await _tres_bases(cliente, esp, "idem")
    lid = await _layout_publicado(cliente, esp, _agrupado(layouts), nome="Grouped")
    execs = {aba: await _executar(cliente, esp, pids[aba]) for aba in ABAS}

    r1 = await _pedir(cliente, ope, lid, execs)
    assert r1.status_code in (200, 202), r1.text
    e1 = r1.json()
    assert (r1.status_code == 200) == (e1["estado"] == "pronto")
    assert e1["tipo"] == "agrupado" and e1["execucao_id"] is None and e1["no_id"] is None
    assert [p["aba"] for p in e1["partes"]] == list(ABAS)
    assert {p["aba"]: p["execucao_id"] for p in e1["partes"]} == execs
    assert all(p["no_id"] == "saida" and p["layout_versao"] == 2 for p in e1["partes"])
    assert e1["linhas_total"] == 2 + 1 + 3
    assert e1["cadeia_id"] is None  # three standalone executions, no shared run

    # A second identical request (different JSON key order) converges on the SAME record
    r2 = await _pedir(cliente, ope, lid, dict(reversed(list(execs.items()))))
    assert r2.status_code in (200, 202) and r2.json()["id"] == e1["id"]

    pronto = await _aguardar_exportacao(cliente, e1["id"], ope)
    assert pronto["estado"] == "pronto", pronto.get("erro")
    assert pronto["linhas_escritas"] == 6 and pronto["bytes"] > 0

    r3 = await _pedir(cliente, ope, lid, execs)
    assert r3.status_code == 200 and r3.json()["id"] == e1["id"]

    # 1 record, 3 parts, 1 artifact, 1 audit entry carrying the parts
    assert (
        await _contar(
            app_engine, "SELECT count(*) FROM relatorio_exportacoes WHERE chave = :c", c=e1["chave"]
        )
        == 1
    )
    async with app_engine.connect() as conn:
        partes = (
            await conn.execute(
                text(
                    "SELECT aba, execucao_id FROM relatorio_exportacao_partes "
                    "WHERE exportacao_id = :x ORDER BY ordem"
                ),
                {"x": uuid.UUID(e1["id"])},
            )
        ).all()
        trilha = (
            (
                await conn.execute(
                    text(
                        "SELECT ator_tipo, depois FROM auditoria WHERE acao = 'exportar' "
                        "AND recurso_tipo = 'relatorio_exportacao' AND recurso_id = :r"
                    ),
                    {"r": uuid.UUID(e1["id"])},
                )
            )
            .mappings()
            .all()
        )
    assert [(a, str(e)) for a, e in partes] == [(aba, execs[aba]) for aba in ABAS]
    assert len(trilha) == 1 and trilha[0]["ator_tipo"] == "usuario"
    assert trilha[0]["depois"]["tipo"] == "agrupado"
    assert [p["aba"] for p in trilha[0]["depois"]["partes"]] == list(ABAS)
    assert _arquivos_no_lake(e1["chave"]) == ["arquivo.xlsx"]

    # download: ONE xlsx with the three sheets, each holding the data of ITS base
    r = await cliente.get(f"/relatorios/exportacoes/{e1['id']}/arquivo", headers=ope)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == list(ABAS)
    assert [c.value for c in wb["SALES"][1]] == ["Sale", "amount"]
    assert list(wb["SALES"].iter_rows(min_row=2, values_only=True)) == [
        ("S1", 10.5),
        ("S2", None),
    ]
    assert wb["SALES"]["B2"].number_format == '"R$ "#,##0.00'
    assert wb["INVOICES"]["B2"].value.date().isoformat() == "2026-08-01"
    assert wb["INVOICES"]["B2"].number_format == "dd/mm/yyyy"
    assert [row[0] for row in wb["PAYMENTS"].iter_rows(min_row=2, values_only=True)] == [
        "P1",
        "P2",
        "P3",
    ]

    # listing by the execution of ONE base includes the grouped export
    lista = (
        await cliente.get(f"/relatorios/exportacoes?execucao_id={execs['INVOICES']}", headers=ope)
    ).json()
    assert [x["id"] for x in lista] == [e1["id"]] and lista[0]["partes"]


async def test_grouped_request_errors_name_the_sheet_and_write_nothing(
    owner_engine, app_engine, cliente
):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.agr.erro@example.com", perfil="especialista"
    )
    pids, layouts = await _tres_bases(cliente, esp, "erro")
    lid = await _layout_publicado(cliente, esp, _agrupado(layouts), nome="Grouped error")
    execs = {aba: await _executar(cliente, esp, pids[aba]) for aba in ("SALES", "PAYMENTS")}
    async with owner_engine.begin() as conn:
        em_voo = (
            await conn.execute(
                text(
                    "INSERT INTO execucoes (pipeline_id, versao_numero, status) "
                    "VALUES (:p, 2, 'executando') RETURNING id"
                ),
                {"p": uuid.UUID(pids["INVOICES"])},
            )
        ).scalar_one()

    r = await _pedir(cliente, esp, lid, {**execs, "INVOICES": str(em_voo)})
    assert r.status_code == 409, r.text
    assert r.json()["erro"]["codigo"] == "execucao_nao_concluida"
    assert "INVOICES" in r.json()["erro"]["mensagem"]

    r = await _pedir(cliente, esp, lid, execs)  # INVOICES missing
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "partes_nao_conferem"
    assert "INVOICES" in r.json()["erro"]["mensagem"]
    r = await _pedir(cliente, esp, lid, {**execs, "INVOICES": str(em_voo), "EXTRA": str(em_voo)})
    assert r.status_code == 400 and "EXTRA" in r.json()["erro"]["mensagem"]

    # execution of ANOTHER pipeline in the sheet: the sheet layout does not match
    r = await _pedir(cliente, esp, lid, {**execs, "INVOICES": execs["SALES"]})
    assert r.status_code == 409 and r.json()["erro"]["codigo"] == "layout_nao_corresponde"
    assert "INVOICES" in r.json()["erro"]["mensagem"]

    # single-source layout requested with parts, and grouped requested with execution + node
    r = await _pedir(cliente, esp, layouts["SALES"], execs)
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "layout_nao_e_agrupado"
    corpo = {"layout_id": lid, "execucao_id": execs["SALES"], "no_id": "saida"}
    r = await cliente.post("/relatorios/exportacoes", json=corpo, headers=esp)
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "layout_agrupado_exige_partes"
    # both modes together, or neither
    r = await cliente.post("/relatorios/exportacoes", json={**corpo, "partes": execs}, headers=esp)
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "pedido_invalido"
    r = await cliente.post("/relatorios/exportacoes", json={"layout_id": lid}, headers=esp)
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "pedido_invalido"

    assert await _contar(app_engine, "SELECT count(*) FROM relatorio_exportacoes") == 0
    assert await _contar(app_engine, "SELECT count(*) FROM relatorio_exportacao_partes") == 0


async def test_publishing_grouped_validates_parts_and_protects_the_base(owner_engine, cliente):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.agr.pub@example.com", perfil="especialista"
    )
    _, ope = await usuario_logado(
        owner_engine, cliente, email="ope.agr.pub@example.com", perfil="operador"
    )
    pids, layouts = await _tres_bases(cliente, esp, "pub")
    lid = (await cliente.post("/relatorios/layouts", json={"nome": "Grp"}, headers=esp)).json()[
        "id"
    ]

    async def _rascunho(definicao, h=esp, base=0):
        return await cliente.put(
            f"/relatorios/layouts/{lid}/rascunho",
            json={"definicao": definicao, "base_numero": base},
            headers=h,
        )

    # sheet pointing at a layout that does not exist / has no published version
    r = await _rascunho(_agrupado({**layouts, "INVOICES": str(uuid.uuid4())}))
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "layout_parte_inexistente"
    rascunho = (
        await cliente.post("/relatorios/layouts", json={"nome": "unpublished"}, headers=esp)
    ).json()["id"]
    r = await _rascunho(_agrupado({**layouts, "INVOICES": rascunho}))
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "layout_parte_nao_publicada"
    # operators cannot author grouped layouts
    r = await _rascunho(_agrupado(layouts), h=ope)
    assert r.status_code == 403

    r = await _rascunho(_agrupado(layouts))
    assert r.status_code == 200, r.text
    r = await cliente.post(
        f"/relatorios/layouts/{lid}/publicar", json={"base_numero": 1}, headers=esp
    )
    assert r.status_code == 200, r.text
    # a grouped layout cannot be a sheet of another grouped layout
    outro = (await cliente.post("/relatorios/layouts", json={"nome": "Agr2"}, headers=esp)).json()[
        "id"
    ]
    r = await cliente.put(
        f"/relatorios/layouts/{outro}/rascunho",
        json={"definicao": _agrupado({**layouts, "INVOICES": lid}), "base_numero": 0},
        headers=esp,
    )
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "layout_agrupado_aninhado"

    # listing by type shows the sheets
    lista = (await cliente.get("/relatorios/layouts?tipo=agrupado", headers=ope)).json()
    assert [x["id"] for x in lista] == [lid]
    assert lista[0]["tipo"] == "agrupado" and lista[0]["pipeline_id"] is None
    assert lista[0]["abas"] == [{"aba": aba, "layout_id": layouts[aba]} for aba in ABAS]
    simples = (await cliente.get("/relatorios/layouts?tipo=simples", headers=ope)).json()
    assert lid not in {x["id"] for x in simples} and layouts["SALES"] in {x["id"] for x in simples}
    r = await cliente.get("/relatorios/layouts?tipo=pivot", headers=ope)
    assert r.status_code == 400

    # deleting the layout of a base that is a sheet of a published grouped layout: 409
    r = await cliente.delete(f"/relatorios/layouts/{layouts['SALES']}", headers=esp)
    assert r.status_code == 409 and r.json()["erro"]["codigo"] == "layout_em_agrupado"


async def test_audit_and_parts_are_atomic_with_the_record(
    owner_engine, app_engine, cliente, monkeypatch
):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.agr.atom@example.com", perfil="especialista"
    )
    pids, layouts = await _tres_bases(cliente, esp, "atom")
    lid = await _layout_publicado(cliente, esp, _agrupado(layouts), nome="Agrupadas atom")
    execs = {aba: await _executar(cliente, esp, pids[aba]) for aba in ABAS}

    async def _quebra(*_a, **_kw):
        raise RuntimeError("audit trail unavailable (intentional)")

    import cflow_core.relatorios.service as service

    monkeypatch.setattr(service, "registrar", _quebra)
    r = await _pedir(cliente, esp, lid, execs)
    assert r.status_code == 500
    assert await _contar(app_engine, "SELECT count(*) FROM relatorio_exportacoes") == 0
    assert await _contar(app_engine, "SELECT count(*) FROM relatorio_exportacao_partes") == 0


# --- eager generation within a run, and failure as state ----------------------------------------


async def _exportacoes_agrupadas(app_engine) -> list[dict]:
    async with app_engine.connect() as conn:
        return [
            dict(r)
            for r in (
                await conn.execute(
                    text(
                        "SELECT id, estado, cadeia_id, gerado_por FROM relatorio_exportacoes "
                        "WHERE tipo = 'agrupado' ORDER BY criado_em"
                    )
                )
            )
            .mappings()
            .all()
        ]


async def test_eager_generates_grouped_when_last_base_of_the_run_completes(
    owner_engine, app_engine, cliente
):
    if not _WORKER["ok"]:
        pytest.skip(f"Test Temporal worker unavailable ({_WORKER['motivo']}).")
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.agr.eager@example.com", perfil="especialista"
    )
    pids, layouts = await _tres_bases(cliente, esp, "eager")
    await _layout_publicado(cliente, esp, _agrupado(layouts, padrao=True), nome="Agr eager")

    # A base running ALONE (outside a run) does not complete the grouped export: siblings are missing.
    await _executar(cliente, esp, pids["SALES"])
    assert await _exportacoes_agrupadas(app_engine) == []

    # The root reads all three bases: the run executes them and the LAST one to finish requests the grouped export.
    raiz = await _pipeline_publicado(
        cliente,
        esp,
        "root-run",
        [
            *(
                {
                    "id": f"src_{aba.lower()}",
                    "tipo": "fonte",
                    "config": {"origem": "pipeline", "pipeline_id": pids[aba], "no_id": "saida"},
                }
                for aba in ABAS
            ),
            {
                "id": "resumo",
                "tipo": "sql",
                "sql": "SELECT (SELECT count(*) FROM src_sales) + (SELECT count(*) FROM "
                "src_invoices) + (SELECT count(*) FROM src_payments) AS linhas",
            },
        ],
        arestas=[{"de": f"src_{aba.lower()}", "para": "resumo"} for aba in ABAS],
    )
    r = await cliente.post(f"/pipelines/{raiz}/executar", headers=esp)
    assert r.status_code == 200, r.text
    res = await aguardar_execucao(cliente, r.json()["id"], esp)
    assert res["status"] == "concluida", res.get("erro")
    async with app_engine.connect() as conn:
        cadeia = (
            await conn.execute(
                text("SELECT cadeia_id FROM execucoes WHERE id = :e"), {"e": uuid.UUID(res["id"])}
            )
        ).scalar_one()
        da_corrida = {
            str(p): str(e)
            for p, e in (
                await conn.execute(
                    text("SELECT pipeline_id, id FROM execucoes WHERE cadeia_id = :c"),
                    {"c": cadeia},
                )
            ).all()
        }
    assert cadeia is not None

    loop = asyncio.get_running_loop()
    fim = loop.time() + 60
    agrupadas: list[dict] = []
    while loop.time() < fim:
        agrupadas = await _exportacoes_agrupadas(app_engine)
        if agrupadas and all(x["estado"] in ("pronto", "falha") for x in agrupadas):
            break
        await asyncio.sleep(0.25)
    assert len(agrupadas) == 1, agrupadas  # ONE, even with two bases completing together
    (ag,) = agrupadas
    assert ag["estado"] == "pronto" and ag["cadeia_id"] == cadeia and ag["gerado_por"] is None
    det = (await cliente.get(f"/relatorios/exportacoes/{ag['id']}", headers=esp)).json()
    assert {p["aba"]: p["execucao_id"] for p in det["partes"]} == {
        aba: da_corrida[pids[aba]] for aba in ABAS
    }
    assert (
        await _contar(
            app_engine,
            "SELECT count(*) FROM auditoria WHERE acao = 'exportar' AND ator_tipo = 'sistema' "
            "AND recurso_id = :r",
            r=ag["id"],
        )
        == 1
    )
    # A request after the eager run: already ready (200), the SAME record.
    r = await _pedir(
        cliente, esp, det["layout_id"], {p["aba"]: p["execucao_id"] for p in det["partes"]}
    )
    assert r.status_code == 200 and r.json()["id"] == str(ag["id"])


async def test_grouped_failure_becomes_state_and_retry_is_a_new_record(
    owner_engine, app_engine, cliente, monkeypatch
):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.agr.falha@example.com", perfil="especialista"
    )
    pids, layouts = await _tres_bases(cliente, esp, "falha")
    lid = await _layout_publicado(cliente, esp, _agrupado(layouts), nome="Agrupadas falha")
    execs = {aba: await _executar(cliente, esp, pids[aba]) for aba in ABAS}

    import cflow_core.relatorios.gerar as gerar_mod

    def _quebra(*_a, **_kw):
        raise RuntimeError("disk full (intentional)")

    monkeypatch.setattr(gerar_mod, "gerar_xlsx_agrupado", _quebra)
    r1 = await _pedir(cliente, esp, lid, execs)
    assert r1.status_code == 202, r1.text
    falha = await _aguardar_exportacao(cliente, r1.json()["id"], esp)
    assert falha["estado"] == "falha" and "RuntimeError" in falha["erro"]
    assert "intentional" not in falha["erro"]
    r = await cliente.get(f"/relatorios/exportacoes/{falha['id']}/arquivo", headers=esp)
    assert r.status_code == 409 and r.json()["erro"]["codigo"] == "exportacao_nao_pronta"

    monkeypatch.undo()
    r2 = await _pedir(cliente, esp, lid, execs)
    assert r2.status_code in (200, 202) and r2.json()["id"] != falha["id"]
    pronto = await _aguardar_exportacao(cliente, r2.json()["id"], esp)
    assert pronto["estado"] == "pronto" and pronto["chave"] == falha["chave"]
    assert (
        await _contar(
            app_engine,
            "SELECT count(*) FROM relatorio_exportacao_partes pt JOIN relatorio_exportacoes x "
            "ON x.id = pt.exportacao_id WHERE x.chave = :c",
            c=falha["chave"],
        )
        == 6  # both attempts keep THEIR parts: the failure is complete evidence
    )


# --- identity and run boundary -------------------------------------------------


async def test_same_layouts_with_other_executions_are_another_artifact(
    owner_engine, app_engine, cliente
):
    """The grouped key carries the execution of EACH sheet. The same layouts over another run, or
    with ONE base from another run, are another record and another file, never the previous
    artifact served again."""
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.agr.d21@example.com", perfil="especialista"
    )
    pids, layouts = await _tres_bases(cliente, esp, "d21")
    lid = await _layout_publicado(cliente, esp, _agrupado(layouts), nome="Agrupadas d21")
    execs1 = {aba: await _executar(cliente, esp, pids[aba]) for aba in ABAS}
    execs2 = {aba: await _executar(cliente, esp, pids[aba]) for aba in ABAS}
    assert set(execs1.values()).isdisjoint(execs2.values())
    pedidos = [execs1, execs2, {**execs1, "INVOICES": execs2["INVOICES"]}]

    registros = []
    for partes in pedidos:
        r = await _pedir(cliente, esp, lid, partes)
        assert r.status_code in (200, 202), r.text
        registros.append(r.json())
    assert len({x["id"] for x in registros}) == 3
    assert len({x["chave"] for x in registros}) == 3
    for x, partes in zip(registros, pedidos, strict=True):
        assert {p["aba"]: p["execucao_id"] for p in x["partes"]} == partes
        pronto = await _aguardar_exportacao(cliente, x["id"], esp)
        assert pronto["estado"] == "pronto", pronto.get("erro")
    assert (
        await _contar(
            app_engine, "SELECT count(*) FROM relatorio_exportacoes WHERE tipo = 'agrupado'"
        )
        == 3
    )
    # And the same request again still converges on ITS record.
    r = await _pedir(cliente, esp, lid, execs2)
    assert r.status_code == 200 and r.json()["id"] == registros[1]["id"]


async def _raiz(cliente, h, nome: str, pids: dict[str, str], abas) -> str:
    """Pipeline that reads the `abas` via `origem: "pipeline"`: triggering it runs those bases in one run."""
    return await _pipeline_publicado(
        cliente,
        h,
        nome,
        [
            *(
                {
                    "id": f"src_{aba.lower()}",
                    "tipo": "fonte",
                    "config": {"origem": "pipeline", "pipeline_id": pids[aba], "no_id": "saida"},
                }
                for aba in abas
            ),
            {
                "id": "resumo",
                "tipo": "sql",
                "sql": "SELECT "
                + " + ".join(f"(SELECT count(*) FROM src_{aba.lower()})" for aba in abas)
                + " AS linhas",
            },
        ],
        arestas=[{"de": f"src_{aba.lower()}", "para": "resumo"} for aba in abas],
    )


async def _corrida(cliente, app_engine, h, raiz: str) -> dict[str, str]:
    """Triggers the root and returns `{pipeline_id: execucao_id}` for the run."""
    r = await cliente.post(f"/pipelines/{raiz}/executar", headers=h)
    assert r.status_code == 200, r.text
    res = await aguardar_execucao(cliente, r.json()["id"], h)
    assert res["status"] == "concluida", res.get("erro")
    async with app_engine.connect() as conn:
        cadeia = (
            await conn.execute(
                text("SELECT cadeia_id FROM execucoes WHERE id = :e"), {"e": uuid.UUID(res["id"])}
            )
        ).scalar_one()
        assert cadeia is not None
        return {
            str(p): str(e)
            for p, e in (
                await conn.execute(
                    text("SELECT pipeline_id, id FROM execucoes WHERE cadeia_id = :c"),
                    {"c": cadeia},
                )
            ).all()
        }


async def test_eager_does_not_complete_grouped_with_base_from_another_run(
    owner_engine, app_engine, cliente
):
    """Run A executes all three bases (and completes the grouped export); run B executes only TWO.
    B's executions must not request the grouped export using A's PAYMENTS: the sibling must come from
    the SAME run."""
    if not _WORKER["ok"]:
        pytest.skip(f"Test Temporal worker unavailable ({_WORKER['motivo']}).")
    from cflow_core.relatorios import service

    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.agr.entre@example.com", perfil="especialista"
    )
    pids, layouts = await _tres_bases(cliente, esp, "entre")
    await _layout_publicado(cliente, esp, _agrupado(layouts, padrao=True), nome="Agr entre")

    corrida_a = await _corrida(
        cliente, app_engine, esp, await _raiz(cliente, esp, "root-a", pids, ABAS)
    )
    assert all(pids[aba] in corrida_a for aba in ABAS)
    loop = asyncio.get_running_loop()
    fim = loop.time() + 60
    while loop.time() < fim:
        agrupadas = await _exportacoes_agrupadas(app_engine)
        if agrupadas and all(x["estado"] in ("pronto", "falha") for x in agrupadas):
            break
        await asyncio.sleep(0.25)
    (ag_a,) = agrupadas
    assert ag_a["estado"] == "pronto"

    so_duas = ("SALES", "INVOICES")
    corrida_b = await _corrida(
        cliente, app_engine, esp, await _raiz(cliente, esp, "root-b", pids, so_duas)
    )
    assert pids["PAYMENTS"] not in corrida_b
    execs_b = [corrida_b[pids[aba]] for aba in so_duas]

    # In B the PAYMENTS base is missing from the run: nothing was requested, even with A's completed.
    assert (
        await _contar(
            app_engine,
            "SELECT count(*) FROM relatorio_exportacao_partes WHERE execucao_id = ANY(:e)",
            e=[uuid.UUID(e) for e in execs_b],
        )
        == 0
    )
    assert [x["id"] for x in await _exportacoes_agrupadas(app_engine)] == [ag_a["id"]]
    async with app_engine.connect() as conn:
        for e in execs_b:
            assert await service.planejar_agrupados(conn, uuid.UUID(e)) == []
        # Control: in A, the base completes the grouped export with A's siblings (not B's, which are
        # more recent), so the plan is not empty for lack of data.
        (plano,) = await service.planejar_agrupados(conn, uuid.UUID(corrida_a[pids["PAYMENTS"]]))
        assert plano["partes"] == {aba: corrida_a[pids[aba]] for aba in ABAS}
