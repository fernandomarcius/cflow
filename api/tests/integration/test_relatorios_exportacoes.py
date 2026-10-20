"""Reports layer through the API and the worker: published layout, completed execution, export as an
idempotent ARTIFACT (2 identical requests = 1 record and 1 file), then download of an xlsx that
openpyxl can open. Also covers: 409 for an unfinished execution, named rejection of legacy export
parameters, RBAC at the edge, audit trail in the SAME transaction, eager generation (smallest
snapshot first), and failures that become state: the next attempt is a new record."""

from __future__ import annotations

import asyncio
import io
import os
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from ._arnes_auth import usuario_logado
from .conftest import _WORKER, aguardar_execucao

pytestmark = pytest.mark.integration

SQL_3_LINHAS = (
    "SELECT * FROM (VALUES "
    "('Alpha', DATE '2026-01-05', 1234.56, 12.34), "
    "('Beta', NULL, NULL, NULL), "
    "('Gamma', DATE '2026-02-10', 20.00, 100.00)"
    ") t(company, sale_date, amount, pct)"
)


def _layout(pipeline_id: str, no_id: str = "n1", **extra) -> dict:
    d = {
        "nome": "Demo",
        "fonte": {"pipeline": pipeline_id, "no": no_id},
        "aba": "Report",
        "colunas": [
            {"coluna": "company", "rotulo": "Company"},
            {"coluna": "sale_date", "rotulo": "Sale date", "formato": "data"},
            {"coluna": "amount", "rotulo": "Amount", "formato": "moeda"},
            {"coluna": "pct", "rotulo": "Pct", "formato": "percentual_pontos"},
        ],
    }
    d.update(extra)
    return d


async def _pipeline_publicado(
    cliente, h, nome: str, nos: list[dict], arestas: list[dict] | None = None
) -> str:
    pid = (await cliente.post("/pipelines", json={"nome": nome}, headers=h)).json()["id"]
    r = await cliente.put(
        f"/pipelines/{pid}/rascunho",
        json={"definicao": {"nos": nos, "arestas": arestas or []}, "base_numero": 0},
        headers=h,
    )
    assert r.status_code == 200, r.text
    r = await cliente.post(
        f"/pipelines/{pid}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=h
    )
    assert r.status_code in (200, 201), r.text
    return pid


async def _layout_publicado(cliente, h, definicao: dict, nome: str = "Layout demo") -> str:
    lid = (await cliente.post("/relatorios/layouts", json={"nome": nome}, headers=h)).json()["id"]
    r = await cliente.put(
        f"/relatorios/layouts/{lid}/rascunho",
        json={"definicao": definicao, "base_numero": 0},
        headers=h,
    )
    assert r.status_code == 200, r.text
    r = await cliente.post(
        f"/relatorios/layouts/{lid}/publicar", json={"base_numero": 1, "nota": "v1"}, headers=h
    )
    assert r.status_code == 200, r.text
    return lid


async def _executar(cliente, h, pid: str) -> str:
    r = await cliente.post(f"/pipelines/{pid}/executar", headers=h)
    assert r.status_code == 200, r.text
    res = await aguardar_execucao(cliente, r.json()["id"], h)
    assert res["status"] == "concluida", res.get("erro")
    return res["id"]


async def _aguardar_exportacao(cliente, eid: str, h, timeout: float = 60.0) -> dict:
    if not _WORKER["ok"]:
        pytest.skip(f"Test Temporal worker unavailable ({_WORKER['motivo']}).")
    loop = asyncio.get_running_loop()
    fim = loop.time() + timeout
    ultimo: dict = {}
    while loop.time() < fim:
        ultimo = (await cliente.get(f"/relatorios/exportacoes/{eid}", headers=h)).json()
        if ultimo.get("estado") in ("pronto", "falha"):
            return ultimo
        await asyncio.sleep(0.25)
    raise AssertionError(f"export {eid} did not finish in {timeout}s ({ultimo.get('estado')})")


async def _contar(engine, sql: str, **params) -> int:
    async with engine.connect() as conn:
        return int((await conn.execute(text(sql), params)).scalar_one())


def _arquivos_no_lake(chave: str) -> list[str]:
    raiz = os.path.join(os.environ["CFLOW_LAKE_DIR"], "relatorios", chave)
    if not os.path.isdir(raiz):
        return []
    return sorted(f for f in os.listdir(raiz) if not f.startswith("."))


# --- end-to-end flow --------------------------------------------------------------------------


async def test_export_is_idempotent_and_xlsx_matches_snapshot(owner_engine, app_engine, cliente):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.rel@example.com", perfil="especialista"
    )
    _, ope = await usuario_logado(
        owner_engine, cliente, email="ope.rel@example.com", perfil="operador"
    )
    pid = await _pipeline_publicado(
        cliente, esp, "rel-demo", [{"id": "n1", "tipo": "sql", "sql": SQL_3_LINHAS}]
    )
    lid = await _layout_publicado(cliente, esp, _layout(pid))
    exec_id = await _executar(cliente, esp, pid)

    corpo = {"execucao_id": exec_id, "no_id": "n1", "layout_id": lid}
    r1 = await cliente.post("/relatorios/exportacoes", json=corpo, headers=ope)
    assert r1.status_code in (200, 202), r1.text
    e1 = r1.json()
    assert e1["estado"] in ("gerando", "pronto")
    assert (r1.status_code == 200) == (e1["estado"] == "pronto")
    assert e1["linhas_total"] == 3 and e1["nome_arquivo"].endswith(".xlsx")

    # a second identical request converges on the SAME record
    r2 = await cliente.post("/relatorios/exportacoes", json=corpo, headers=ope)
    assert r2.status_code in (200, 202)
    assert r2.json()["id"] == e1["id"]

    pronto = await _aguardar_exportacao(cliente, e1["id"], ope)
    assert pronto["estado"] == "pronto", pronto.get("erro")
    assert pronto["linhas_escritas"] == 3 and pronto["bytes"] > 0 and pronto["concluido_em"]

    # third request, already ready: 200 and no new generation
    r3 = await cliente.post("/relatorios/exportacoes", json=corpo, headers=ope)
    assert r3.status_code == 200 and r3.json()["id"] == e1["id"]

    # 1 record, 1 artifact, 1 audit entry
    assert (
        await _contar(
            app_engine, "SELECT count(*) FROM relatorio_exportacoes WHERE chave = :c", c=e1["chave"]
        )
        == 1
    )
    assert _arquivos_no_lake(e1["chave"]) == ["arquivo.xlsx"]
    assert (
        await _contar(
            app_engine,
            "SELECT count(*) FROM auditoria WHERE acao = 'exportar' "
            "AND recurso_tipo = 'relatorio_exportacao' AND recurso_id = :r",
            r=uuid.UUID(e1["id"]),
        )
        == 1
    )

    # download: an xlsx openpyxl can open, with the snapshot values
    r = await cliente.get(f"/relatorios/exportacoes/{e1['id']}/arquivo", headers=ope)
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert e1["nome_arquivo"] in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Report", "Carimbo"]
    ws = wb["Report"]
    assert [c.value for c in ws[1]] == ["Company", "Sale date", "Amount", "Pct"]
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 3
    assert linhas[0][0] == "Alpha" and linhas[0][1].date().isoformat() == "2026-01-05"
    assert float(linhas[0][2]) == 1234.56 and float(linhas[0][3]) == 12.34
    assert linhas[1] == ("Beta", None, None, None)  # NULL comes out empty, never 0
    assert float(linhas[2][3]) == 100.0
    assert ws["D2"].number_format == '0.00"%"' and ws["B2"].number_format == "dd/mm/yyyy"
    carimbo = {r[0]: r[1] for r in wb["Carimbo"].iter_rows(min_row=2, values_only=True)}
    assert carimbo["execucao_id"] == exec_id and carimbo["no_id"] == "n1"
    assert carimbo["layout_id"] == lid and carimbo["layout_versao"] == 2
    assert carimbo["gerado_por"] == "Teste" and carimbo["gerado_em"].endswith("-03:00")
    assert carimbo["hash_snapshot"]

    # listing by execution (history) and GET by id
    lista = (
        await cliente.get(f"/relatorios/exportacoes?execucao_id={exec_id}", headers=ope)
    ).json()
    assert [x["id"] for x in lista] == [e1["id"]]
    det = (await cliente.get(f"/relatorios/exportacoes/{e1['id']}", headers=ope)).json()
    assert det["estado"] == "pronto" and det["gerado_por"] == "Teste"


async def test_unfinished_execution_returns_409(owner_engine, app_engine, cliente):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp409@example.com", perfil="especialista"
    )
    pid = await _pipeline_publicado(
        cliente, esp, "rel-409", [{"id": "n1", "tipo": "sql", "sql": SQL_3_LINHAS}]
    )
    lid = await _layout_publicado(cliente, esp, _layout(pid))
    async with owner_engine.begin() as conn:
        exec_id = (
            await conn.execute(
                text(
                    "INSERT INTO execucoes (pipeline_id, versao_numero, status) "
                    "VALUES (:p, 2, 'executando') RETURNING id"
                ),
                {"p": uuid.UUID(pid)},
            )
        ).scalar_one()
    r = await cliente.post(
        "/relatorios/exportacoes",
        json={"execucao_id": str(exec_id), "no_id": "n1", "layout_id": lid},
        headers=esp,
    )
    assert r.status_code == 409, r.text
    assert r.json()["erro"]["codigo"] == "execucao_nao_concluida"
    assert await _contar(app_engine, "SELECT count(*) FROM relatorio_exportacoes") == 0


@pytest.mark.parametrize("extra", ["projectId", "amount", "operator", "end_date", "recordDate"])
async def test_legacy_export_parameter_is_rejected_by_name(owner_engine, cliente, extra):
    _, ope = await usuario_logado(
        owner_engine, cliente, email=f"ope.{extra}@example.com", perfil="operador"
    )
    r = await cliente.post(
        "/relatorios/exportacoes",
        json={
            "execucao_id": str(uuid.uuid4()),
            "no_id": "n1",
            "layout_id": str(uuid.uuid4()),
            extra: "1",
        },
        headers=ope,
    )
    assert r.status_code == 400, r.text
    assert r.json()["erro"]["codigo"] == "parametro_nao_e_do_relatorio"
    assert extra in r.json()["erro"]["mensagem"]


# --- RBAC at the edge --------------------------------------------------------------------------------


async def test_rbac_operator_views_and_downloads_but_cannot_author_layout(
    owner_engine, app_engine, cliente
):
    _, ope = await usuario_logado(
        owner_engine, cliente, email="ope.rbac@example.com", perfil="operador"
    )
    _, adm = await usuario_logado(
        owner_engine, cliente, email="adm.rbac@example.com", perfil="admin"
    )
    # operators cannot create or publish layouts
    r = await cliente.post("/relatorios/layouts", json={"nome": "x"}, headers=ope)
    assert r.status_code == 403 and r.json()["erro"]["codigo"] == "nao_autorizado"
    # admin creates; operators cannot edit
    lid = (await cliente.post("/relatorios/layouts", json={"nome": "x"}, headers=adm)).json()["id"]
    pid = await _pipeline_publicado(
        cliente, adm, "rel-rbac", [{"id": "n1", "tipo": "sql", "sql": SQL_3_LINHAS}]
    )
    r = await cliente.put(
        f"/relatorios/layouts/{lid}/rascunho",
        json={"definicao": _layout(pid), "base_numero": 0},
        headers=ope,
    )
    assert r.status_code == 403
    r = await cliente.post(
        f"/relatorios/layouts/{lid}/publicar", json={"base_numero": 0}, headers=ope
    )
    assert r.status_code == 403
    r = await cliente.delete(f"/relatorios/layouts/{lid}", headers=ope)
    assert r.status_code == 403
    # operators can read layouts
    assert (await cliente.get("/relatorios/layouts", headers=ope)).status_code == 200
    # no session: 401 everywhere
    for metodo, rota in [
        ("post", "/relatorios/exportacoes"),
        ("get", "/relatorios/exportacoes"),
        ("get", f"/relatorios/exportacoes/{uuid.uuid4()}"),
        ("get", f"/relatorios/exportacoes/{uuid.uuid4()}/arquivo"),
        ("post", "/relatorios/layouts"),
        ("get", "/relatorios/layouts"),
    ]:
        r = await (cliente.post(rota, json={}) if metodo == "post" else cliente.get(rota))
        assert r.status_code == 401, (metodo, rota, r.text)
        assert r.json()["erro"]["codigo"] == "nao_autenticado"


# --- audit trail in the SAME transaction ---------------------------------------------------------------


async def test_layout_audit_trail_create_and_publish(owner_engine, app_engine, cliente):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.trilha@example.com", perfil="especialista"
    )
    pid = await _pipeline_publicado(
        cliente, esp, "rel-trilha", [{"id": "n1", "tipo": "sql", "sql": SQL_3_LINHAS}]
    )
    lid = await _layout_publicado(cliente, esp, _layout(pid), nome="Trilha")
    async with app_engine.connect() as conn:
        rows = (
            (
                await conn.execute(
                    text(
                        "SELECT acao, ator_tipo, depois FROM auditoria "
                        "WHERE recurso_tipo = 'relatorio_layout' AND recurso_id = :r ORDER BY id"
                    ),
                    {"r": uuid.UUID(lid)},
                )
            )
            .mappings()
            .all()
        )
    assert [r["acao"] for r in rows] == ["criar", "editar", "publicar"]
    assert all(r["ator_tipo"] == "usuario" for r in rows)
    assert rows[0]["depois"] == {"nome": "Trilha"}
    assert rows[2]["depois"]["numero"] == 2 and rows[2]["depois"]["nota"] == "v1"


async def test_audit_trail_and_write_are_atomic(owner_engine, app_engine, cliente, monkeypatch):
    """If the audit trail cannot be written, the business write does not persist either (layout and export)."""
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.atomico@example.com", perfil="especialista"
    )
    pid = await _pipeline_publicado(
        cliente, esp, "rel-atomico", [{"id": "n1", "tipo": "sql", "sql": SQL_3_LINHAS}]
    )
    lid = await _layout_publicado(cliente, esp, _layout(pid))
    exec_id = await _executar(cliente, esp, pid)

    async def _quebra(*_a, **_kw):
        raise RuntimeError("audit trail unavailable (intentional)")

    import cflow_core.modelagem.versionado as versionado
    import cflow_core.relatorios.service as service

    antes_layouts = await _contar(app_engine, "SELECT count(*) FROM relatorio_layouts")
    monkeypatch.setattr(versionado, "registrar", _quebra)
    r = await cliente.post("/relatorios/layouts", json={"nome": "no trail"}, headers=esp)
    assert r.status_code == 500 and r.json()["erro"]["codigo"] == "erro_interno"
    assert await _contar(app_engine, "SELECT count(*) FROM relatorio_layouts") == antes_layouts

    monkeypatch.setattr(service, "registrar", _quebra)
    r = await cliente.post(
        "/relatorios/exportacoes",
        json={"execucao_id": exec_id, "no_id": "n1", "layout_id": lid},
        headers=esp,
    )
    assert r.status_code == 500
    assert await _contar(app_engine, "SELECT count(*) FROM relatorio_exportacoes") == 0
    assert (
        await _contar(
            app_engine,
            "SELECT count(*) FROM auditoria WHERE recurso_tipo = 'relatorio_exportacao'",
        )
        == 0
    )


# --- layout validation against the snapshot -----------------------------------------------------


async def test_publishing_layout_with_column_outside_snapshot_fails(owner_engine, cliente):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.col@example.com", perfil="especialista"
    )
    pid = await _pipeline_publicado(
        cliente, esp, "rel-col", [{"id": "n1", "tipo": "sql", "sql": SQL_3_LINHAS}]
    )
    await _executar(cliente, esp, pid)
    lid = (await cliente.post("/relatorios/layouts", json={"nome": "col"}, headers=esp)).json()[
        "id"
    ]
    d = _layout(pid)
    d["colunas"].append({"coluna": "missing_col", "rotulo": "Missing"})
    r = await cliente.put(
        f"/relatorios/layouts/{lid}/rascunho", json={"definicao": d, "base_numero": 0}, headers=esp
    )
    assert r.status_code == 200, r.text  # draft accepts it; publishing is what checks the snapshot
    r = await cliente.post(
        f"/relatorios/layouts/{lid}/publicar", json={"base_numero": 1}, headers=esp
    )
    assert r.status_code == 400, r.text
    assert r.json()["erro"]["codigo"] == "layout_coluna_inexistente"
    assert "'missing_col'" in r.json()["erro"]["mensagem"]
    # format outside the vocabulary and data transformation: rejected already at draft time
    d = _layout(pid)
    d["colunas"][2]["formato"] = "moeda_dividida_por_100"
    r = await cliente.put(
        f"/relatorios/layouts/{lid}/rascunho", json={"definicao": d, "base_numero": 1}, headers=esp
    )
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "layout_formato_invalido"
    d = _layout(pid, total="soma")
    r = await cliente.put(
        f"/relatorios/layouts/{lid}/rascunho", json={"definicao": d, "base_numero": 1}, headers=esp
    )
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "layout_transforma_dado"
    # node that does not exist in the pipeline
    r = await cliente.put(
        f"/relatorios/layouts/{lid}/rascunho",
        json={"definicao": _layout(pid, no_id="nao_existe"), "base_numero": 1},
        headers=esp,
    )
    assert r.status_code == 400 and r.json()["erro"]["codigo"] == "layout_no_inexistente"


async def test_export_with_layout_of_another_node_returns_409(owner_engine, cliente):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.outro@example.com", perfil="especialista"
    )
    pid = await _pipeline_publicado(
        cliente,
        esp,
        "rel-outro",
        [
            {"id": "n1", "tipo": "sql", "sql": SQL_3_LINHAS},
            {"id": "n2", "tipo": "sql", "sql": "SELECT 1 AS one"},
        ],
        arestas=[{"de": "n1", "para": "n2"}],
    )
    lid = await _layout_publicado(cliente, esp, _layout(pid, no_id="n1"))
    exec_id = await _executar(cliente, esp, pid)
    r = await cliente.post(
        "/relatorios/exportacoes",
        json={"execucao_id": exec_id, "no_id": "n2", "layout_id": lid},
        headers=esp,
    )
    assert r.status_code == 409 and r.json()["erro"]["codigo"] == "layout_nao_corresponde"


# --- eager generation and failure as state ----------------------------------------------


async def test_eager_enqueues_default_layouts_smallest_snapshot_first(
    owner_engine, app_engine, cliente
):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.eager@example.com", perfil="especialista"
    )
    pid = await _pipeline_publicado(
        cliente,
        esp,
        "rel-eager",
        [
            {
                "id": "grande",
                "tipo": "sql",
                "sql": "SELECT i AS n, 'x' || i AS txt FROM range(2000) t(i)",
            },
            {"id": "pequeno", "tipo": "sql", "sql": "SELECT 1 AS n, 'a' AS txt"},
            {"id": "sem_layout", "tipo": "sql", "sql": "SELECT 2 AS n"},
        ],
        arestas=[{"de": "grande", "para": "pequeno"}, {"de": "pequeno", "para": "sem_layout"}],
    )
    colunas = [{"coluna": "n", "formato": "inteiro"}, {"coluna": "txt"}]
    # published BEFORE any execution exists: the warning says validation is deferred to export
    for no in ("grande", "pequeno"):
        lid = (
            await cliente.post("/relatorios/layouts", json={"nome": f"default {no}"}, headers=esp)
        ).json()["id"]
        await cliente.put(
            f"/relatorios/layouts/{lid}/rascunho",
            json={
                "definicao": _layout(pid, no_id=no, colunas=colunas, padrao=True),
                "base_numero": 0,
            },
            headers=esp,
        )
        r = await cliente.post(
            f"/relatorios/layouts/{lid}/publicar", json={"base_numero": 1}, headers=esp
        )
        assert r.status_code == 200 and r.json()["conferido"] is False and r.json()["aviso"]

    exec_id = await _executar(cliente, esp, pid)

    # the plan (order) is pure and queryable: smallest snapshot first
    from cflow_core.relatorios import service

    async with app_engine.connect() as conn:
        plano = await service.planejar_padroes(conn, uuid.UUID(exec_id))
    assert [p["no_id"] for p in plano] == ["pequeno", "grande"]
    assert plano[0]["bytes_snapshot"] < plano[1]["bytes_snapshot"]

    # both exports were created on their own when the run completed (gerado_por is None) and become ready
    loop = asyncio.get_running_loop()
    fim = loop.time() + 60
    lista: list[dict] = []
    while loop.time() < fim:
        lista = (
            await cliente.get(f"/relatorios/exportacoes?execucao_id={exec_id}", headers=esp)
        ).json()
        if len(lista) == 2 and all(x["estado"] == "pronto" for x in lista):
            break
        await asyncio.sleep(0.25)
    assert len(lista) == 2 and all(x["estado"] == "pronto" for x in lista), lista
    assert {x["no_id"] for x in lista} == {"pequeno", "grande"}
    assert all(x["gerado_por"] is None for x in lista)
    por_no = {x["no_id"]: x for x in lista}
    assert por_no["pequeno"]["criado_em"] <= por_no["grande"]["criado_em"]
    assert por_no["grande"]["linhas_total"] == 2000

    async with app_engine.connect() as conn:
        atores = (
            (
                await conn.execute(
                    text(
                        "SELECT DISTINCT ator_tipo FROM auditoria WHERE acao = 'exportar' "
                        "AND recurso_tipo = 'relatorio_exportacao'"
                    )
                )
            )
            .scalars()
            .all()
        )
        resumo = await service.resumo_da_execucao(conn, uuid.UUID(exec_id))
    assert atores == ["sistema"]
    assert resumo == {"prontos": 2, "total": 2, "falhas": 0}

    # request after the eager run: already ready (200), same record
    r = await cliente.post(
        "/relatorios/exportacoes",
        json={
            "execucao_id": exec_id,
            "no_id": "pequeno",
            "layout_id": por_no["pequeno"]["layout_id"],
        },
        headers=esp,
    )
    assert r.status_code == 200 and r.json()["id"] == por_no["pequeno"]["id"]


async def test_failure_becomes_state_and_retry_is_a_new_record(
    owner_engine, app_engine, cliente, monkeypatch
):
    _, esp = await usuario_logado(
        owner_engine, cliente, email="esp.falha@example.com", perfil="especialista"
    )
    pid = await _pipeline_publicado(
        cliente, esp, "rel-falha", [{"id": "n1", "tipo": "sql", "sql": SQL_3_LINHAS}]
    )
    lid = await _layout_publicado(cliente, esp, _layout(pid))
    exec_id = await _executar(cliente, esp, pid)

    import cflow_core.relatorios.gerar as gerar_mod

    def _quebra(*_a, **_kw):
        raise RuntimeError("disk full (intentional)")

    # The test worker runs in the SAME process: patching the module attribute reaches the activity.
    monkeypatch.setattr(gerar_mod, "gerar_xlsx", _quebra)
    corpo = {"execucao_id": exec_id, "no_id": "n1", "layout_id": lid}
    r1 = await cliente.post("/relatorios/exportacoes", json=corpo, headers=esp)
    assert r1.status_code == 202, r1.text
    falha = await _aguardar_exportacao(cliente, r1.json()["id"], esp)
    assert falha["estado"] == "falha"
    assert "RuntimeError" in falha["erro"]
    assert (
        "intentional" not in falha["erro"]
    )  # internal detail does not leak; the log has the stack
    r = await cliente.get(f"/relatorios/exportacoes/{falha['id']}/arquivo", headers=esp)
    assert r.status_code == 409 and r.json()["erro"]["codigo"] == "exportacao_nao_pronta"

    monkeypatch.undo()
    r2 = await cliente.post("/relatorios/exportacoes", json=corpo, headers=esp)
    assert r2.status_code in (200, 202), r2.text
    assert r2.json()["id"] != falha["id"]  # NEW record: the failure stays as evidence
    pronto = await _aguardar_exportacao(cliente, r2.json()["id"], esp)
    assert pronto["estado"] == "pronto" and pronto["chave"] == falha["chave"]
    assert (
        await _contar(
            app_engine,
            "SELECT count(*) FROM relatorio_exportacoes WHERE chave = :c",
            c=falha["chave"],
        )
        == 2
    )
    assert _arquivos_no_lake(falha["chave"]) == ["arquivo.xlsx"]
