"""`{{coluna.x}}` in the layout preamble: the value the NODE computed, read from the snapshot.

It exists for sheets whose title and loose cells show numbers that are not columns of the table
(for example an average). As a snapshot column, the layout only writes into a cell what the
pipeline already computed -- the layer still calculates nothing.

The refusals are what matter here: a missing column must not become an empty cell that looks
delivered, and a column with two values must not become "the layout picked one".
"""

from __future__ import annotations

import uuid

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from cflow_core.errors import ErroAplicacao
from cflow_core.relatorios import gerar, layout


def _definicao(**extra) -> dict:
    return {
        "fonte": {"pipeline": str(uuid.uuid4()), "no": "saida"},
        "colunas": [{"coluna": "mes", "rotulo": ""}, {"coluna": "valor", "rotulo": "AMOUNT"}],
        "titulo": {"texto": "REPORT - {{coluna.project_name}} - {{param.period_start:%m_%Y}}"},
        "linhas_fixas": [
            {"linha": 2, "coluna": 1, "texto": "average rate"},
            {"linha": 3, "coluna": 1, "valor": "{{coluna.avg_rate}}", "escala": 0.01},
        ],
        "primeira_linha_dado": 5,
        "linha_cabecalho": 4,
        "carimbo": False,
        **extra,
    }


def _snapshot(tmp_path, **colunas) -> str:
    dados = {"mes": ["09/2026", "10/2026"], "valor": [1000.0, 2000.0], **colunas}
    caminho = tmp_path / "snap.parquet"
    pq.write_table(pa.table(dados), caminho)
    return str(caminho)


def _gerar(tmp_path, snapshot: str, **kw):
    destino = tmp_path / "r.xlsx"
    gerar.gerar_xlsx(snapshot, _definicao(), destino, **kw)
    return load_workbook(destino).worksheets[0]


def test_colunas_do_preambulo_na_ordem_e_sem_repetir() -> None:
    d = _definicao()
    d["linhas_fixas"].append({"linha": 1, "coluna": 2, "texto": "{{coluna.avg_rate}}"})
    assert layout.colunas_do_preambulo(d) == ["project_name", "avg_rate"]


def test_titulo_e_linha_fixa_leem_a_coluna_constante_do_snapshot(tmp_path) -> None:
    snap = _snapshot(tmp_path, project_name=["Project X", "Project X"], avg_rate=[47.41, 47.41])
    ws = _gerar(tmp_path, snap, parametros={"period_start": "2026-08-01"})
    assert ws.cell(row=1, column=1).value == "REPORT - Project X - 08_2026"
    # The scale is presentation: 47.41 points in a percentage cell is 0.4741.
    assert ws.cell(row=3, column=1).value == pytest.approx(0.4741)


def test_coluna_do_preambulo_ausente_no_snapshot_e_recusada_nomeando_a_coluna(tmp_path) -> None:
    """An empty cell that looks delivered is the defect `layout_coluna_inexistente` exists to
    prevent -- it applies to the preamble cell as it does to the table column."""
    with pytest.raises(ErroAplicacao) as exc:
        layout.validar_contra_snapshot(_definicao(), ["mes", "valor", "project_name"])
    assert exc.value.codigo == "layout_coluna_inexistente"
    assert "avg_rate" in exc.value.mensagem

    snap = _snapshot(tmp_path, project_name=["Project X", "Project X"])
    with pytest.raises(ErroAplicacao) as exc2:
        _gerar(tmp_path, snap)
    assert exc2.value.codigo == "layout_coluna_inexistente"


def test_coluna_com_dois_valores_e_recusada__o_layout_nao_escolhe(tmp_path) -> None:
    snap = _snapshot(tmp_path, project_name=["Project X", "Project X"], avg_rate=[47.41, 33.68])
    with pytest.raises(ErroAplicacao) as exc:
        _gerar(tmp_path, snap)
    assert exc.value.codigo == "layout_valor_nao_constante"
    assert "avg_rate" in exc.value.mensagem


def test_coluna_so_com_nulo_resolve_vazio__nunca_zero(tmp_path) -> None:
    snap = _snapshot(
        tmp_path,
        project_name=pa.array([None, None], pa.string()),
        avg_rate=pa.array([None, None], pa.float64()),
    )
    ws = _gerar(tmp_path, snap)
    assert ws.cell(row=1, column=1).value == "REPORT -  - "
    assert ws.cell(row=3, column=1).value in (None, "")
