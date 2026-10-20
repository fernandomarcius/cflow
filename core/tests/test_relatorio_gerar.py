"""Node Parquet -> xlsx via the layout. What is proven: the cell value is the Parquet's (no /100,
no rounding), NULL becomes an empty cell (never 0), a points-percentage column comes out as
`0.00"%"`, a date comes out as a date, the Carimbo sheet says where the file came from, and above
the cap the data continues on the next sheet."""

from __future__ import annotations

import uuid
from datetime import date, datetime
from decimal import Decimal

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from cflow_core.errors import ErroAplicacao
from cflow_core.relatorios import gerar, layout

PIPELINE = str(uuid.uuid4())


def _parquet(caminho, n: int = 3):
    companies = [f"Company {i}" for i in range(n)]
    datas = [date(2026, 1, 1 + (i % 27)) for i in range(n)]
    valores = [Decimal("1234.56") + i for i in range(n)]
    pct = [12.34 + i for i in range(n)]
    qtd = [10 + i for i in range(n)]
    if n >= 2:  # the 2nd row has NULL in every typed column -- the case a naive writer turns into 0
        datas[1] = None
        valores[1] = None
        pct[1] = None
        qtd[1] = None
    tabela = pa.table(
        {
            "Company": pa.array(companies, pa.string()),
            "Sale_Date": pa.array(datas, pa.date32()),
            "Amount": pa.array(valores, pa.decimal128(18, 2)),
            "Rate_Pct": pa.array(pct, pa.float64()),
            "Qty": pa.array(qtd, pa.int64()),
            "column_outside_layout": pa.array(["x"] * n, pa.string()),
        }
    )
    pq.write_table(tabela, caminho)
    return tabela


def _layout(**extra) -> dict:
    d = {
        "nome": "Report 1 - Net present value",
        "fonte": {"pipeline": PIPELINE, "no": "full_portfolio_node"},
        "aba": "REPORT 1. NET PRESENT VALUE",
        "colunas": [
            {"coluna": "Company", "rotulo": "Company"},
            {"coluna": "Sale_Date", "rotulo": "Sale Date", "formato": "data"},
            {"coluna": "Amount", "rotulo": "Value", "formato": "moeda"},
            {"coluna": "Rate_Pct", "rotulo": "Rate %", "formato": "percentual_pontos"},
            {"coluna": "Qty", "rotulo": "Qty", "formato": "inteiro"},
        ],
    }
    d.update(extra)
    return d


CARIMBO = [
    ("execucao_id", "11111111-2222-3333-4444-555555555555"),
    ("no_id", "full_portfolio_node"),
    ("layout", "Report 1 - Net present value"),
    ("layout_versao", 3),
    ("gerado_por", "Alice"),
    ("gerado_em", "2026-09-05T10:00:00-03:00"),
    ("hash_snapshot", "abc123"),
]


def test_valor_da_celula_e_o_do_parquet_e_null_sai_vazio(tmp_path):
    parquet = tmp_path / "n.parquet"
    tabela = _parquet(parquet)
    saida = tmp_path / "saida.xlsx"
    res = gerar.gerar_xlsx(parquet, _layout(), saida, carimbo=CARIMBO)
    assert res["linhas"] == 3 and res["bytes"] > 0
    assert res["abas"] == ["REPORT 1. NET PRESENT VALUE", "Carimbo"]

    wb = load_workbook(saida)
    ws = wb["REPORT 1. NET PRESENT VALUE"]
    assert [c.value for c in ws[1]] == ["Company", "Sale Date", "Value", "Rate %", "Qty"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:E4"
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 3
    # 1st row: identical to the Parquet (a number is a number, no /100; Decimal becomes the same number)
    assert linhas[0][0] == "Company 0"
    assert linhas[0][1] == datetime(2026, 1, 1)
    assert Decimal(str(linhas[0][2])) == tabela.column("Amount")[0].as_py()
    assert linhas[0][3] == 12.34  # percentage points, as in the node -- NOT 0.1234
    assert linhas[0][4] == 10
    # 2nd row: NULL is an empty cell, never 0
    assert linhas[1] == ("Company 1", None, None, None, None)
    # a column outside the layout is not written
    assert ws.max_column == 5
    # the format belongs to the CELL
    formatos = [c.number_format for c in ws[2]]
    assert formatos[1] == "dd/mm/yyyy"
    assert formatos[2] == '"R$ "#,##0.00'
    assert formatos[3] == '0.00"%"'
    assert formatos[4] == "0"
    assert ws["B2"].is_date


def test_aba_carimbo_diz_de_onde_o_arquivo_veio(tmp_path):
    parquet = tmp_path / "n.parquet"
    _parquet(parquet)
    saida = tmp_path / "saida.xlsx"
    gerar.gerar_xlsx(parquet, _layout(), saida, carimbo=CARIMBO)
    ws = load_workbook(saida)["Carimbo"]
    pares = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert pares["execucao_id"] == "11111111-2222-3333-4444-555555555555"
    assert pares["no_id"] == "full_portfolio_node"
    assert pares["layout"] == "Report 1 - Net present value" and pares["layout_versao"] == 3
    assert pares["gerado_por"] == "Alice"
    assert pares["gerado_em"] == "2026-09-05T10:00:00-03:00"
    assert pares["hash_snapshot"] == "abc123"


def test_sem_carimbo_quando_o_layout_desliga(tmp_path):
    parquet = tmp_path / "n.parquet"
    _parquet(parquet)
    saida = tmp_path / "saida.xlsx"
    res = gerar.gerar_xlsx(parquet, _layout(carimbo=False), saida, carimbo=CARIMBO)
    assert res["abas"] == ["REPORT 1. NET PRESENT VALUE"]


def test_acima_do_teto_o_dado_continua_na_aba_seguinte(tmp_path):
    parquet = tmp_path / "n.parquet"
    _parquet(parquet, n=7)
    saida = tmp_path / "saida.xlsx"
    progresso: list[int] = []
    res = gerar.gerar_xlsx(
        parquet,
        _layout(aba="Base"),
        saida,
        carimbo=CARIMBO,
        teto_linhas=3,
        progresso=progresso.append,
    )
    assert res["linhas"] == 7
    assert res["abas"] == ["Base (parte 1)", "Base (parte 2)", "Base (parte 3)", "Carimbo"]
    assert progresso[-1] == 7
    wb = load_workbook(saida)
    total = 0
    for nome in res["abas"][:-1]:
        ws = wb[nome]
        assert [c.value for c in ws[1]] == ["Company", "Sale Date", "Value", "Rate %", "Qty"]
        dados = list(ws.iter_rows(min_row=2, values_only=True))
        assert 1 <= len(dados) <= 3
        total += len(dados)
    assert total == 7
    # order preserved across parts
    assert wb["Base (parte 1)"]["A2"].value == "Company 0"
    assert wb["Base (parte 2)"]["A2"].value == "Company 3"
    assert wb["Base (parte 3)"]["A2"].value == "Company 6"


def test_no_teto_exato_nao_abre_parte_vazia(tmp_path):
    parquet = tmp_path / "n.parquet"
    _parquet(parquet, n=6)
    saida = tmp_path / "saida.xlsx"
    res = gerar.gerar_xlsx(parquet, _layout(aba="Base", carimbo=False), saida, teto_linhas=3)
    assert res["abas"] == ["Base (parte 1)", "Base (parte 2)"]


def test_teto_padrao_e_o_do_excel():
    assert layout.TETO_LINHAS_EXCEL == 1_048_576
    assert gerar.TETO_LINHAS_DADO == 1_048_575


def test_coluna_inexistente_no_parquet_e_erro_antes_de_escrever(tmp_path):
    parquet = tmp_path / "n.parquet"
    _parquet(parquet)
    d = _layout()
    d["colunas"].append({"coluna": "Does_Not_Exist"})
    with pytest.raises(ErroAplicacao) as exc:
        gerar.gerar_xlsx(parquet, d, tmp_path / "x.xlsx")
    assert exc.value.codigo == "layout_coluna_inexistente"
    assert not (tmp_path / "x.xlsx").exists()


def test_parquet_vazio_gera_so_cabecalho(tmp_path):
    parquet = tmp_path / "n.parquet"
    _parquet(parquet, n=0)
    saida = tmp_path / "saida.xlsx"
    res = gerar.gerar_xlsx(parquet, _layout(aba="Vazia", carimbo=False), saida)
    assert res["linhas"] == 0 and res["abas"] == ["Vazia"]
    ws = load_workbook(saida)["Vazia"]
    assert ws.max_row == 1


def test_celula_nao_inventa_valor():
    assert gerar._celula(None) is None
    assert gerar._celula(float("nan")) is None
    assert gerar._celula(Decimal("0.10")) == Decimal("0.10")
    assert gerar._celula(0) == 0
    assert gerar._celula(b"\x01\x02") == "0102"
    assert gerar._celula([1, "a"]) == '[1, "a"]'


# --- measured limits of the format ---------------------------------------------------------------


def test_double_de_17_digitos_perde_o_ultimo__limite_do_xlsx_nao_da_camada(tmp_path):
    """The xlsx stores numbers as decimal text and openpyxl serializes with 16 significant digits:
    `1242.8500000000001` (17) comes out as `1242.85`. There is no way to preserve it -- Decimal goes
    through the SAME serialization, and Excel only reads 15 digits.

    The fix belongs at the SOURCE: a node that sums money emits `DECIMAL`, and then the cell is
    bit-exact (the second half of this test proves it). The layer rounds nothing on its own.
    """
    parquet = tmp_path / "n.parquet"
    tabela = pa.table(
        {
            "dirty": pa.array([1242.8500000000001, 426.43999999999994], pa.float64()),
            "clean": pa.array([Decimal("1242.85"), Decimal("426.44")], pa.decimal128(18, 2)),
        }
    )
    pq.write_table(tabela, parquet)
    d = {
        "nome": "limite",
        "fonte": {"pipeline": PIPELINE, "no": "n"},
        "aba": "L",
        "colunas": [
            {"coluna": "dirty", "formato": "moeda"},
            {"coluna": "clean", "formato": "moeda"},
        ],
        "carimbo": False,
    }
    saida = tmp_path / "s.xlsx"
    gerar.gerar_xlsx(parquet, d, saida)
    ws = load_workbook(saida)["L"]
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    # 17-digit DOUBLE: what fits the format, and nothing beyond it
    assert linhas[0][0] == float(f"{1242.8500000000001:.16g}") == 1242.85
    assert linhas[1][0] == float(f"{426.43999999999994:.16g}")
    # DECIMAL with a money scale: bit-exact, the path to a zero difference
    assert Decimal(str(linhas[0][1])) == Decimal("1242.85")
    assert Decimal(str(linhas[1][1])) == Decimal("426.44")


def test_inteiro_acima_de_2_elevado_53_nao_sobrevive_como_numero(tmp_path):
    """A long identifier (16+ digits) becomes scientific notation and LOSES a digit in the xlsx.
    This is a regression guard, not an open defect: a long-identifier column must come out as TEXT
    from the node, never as an integer."""
    parquet = tmp_path / "n.parquet"
    pq.write_table(
        pa.table(
            {
                "id_num": pa.array([12345678901234567], pa.int64()),
                "id_txt": pa.array(["12345678901234567"], pa.string()),
            }
        ),
        parquet,
    )
    d = {
        "nome": "id",
        "fonte": {"pipeline": PIPELINE, "no": "n"},
        "aba": "I",
        "colunas": [{"coluna": "id_num"}, {"coluna": "id_txt", "formato": "texto"}],
        "carimbo": False,
    }
    saida = tmp_path / "s.xlsx"
    gerar.gerar_xlsx(parquet, d, saida)
    ws = load_workbook(saida)["I"]
    numero, texto = next(ws.iter_rows(min_row=2, values_only=True))
    assert numero != 12345678901234567  # corrupted by the format
    assert texto == "12345678901234567"  # intact


def test_celula_vazia_no_FIM_da_linha_nao_e_escrita__quem_le_tem_de_completar(tmp_path):
    """openpyxl in write-only mode omits an empty cell at the END of a row: the row comes back
    SHORTER than the header. The file is valid (Excel reads an absent cell as empty), but anyone
    comparing cell by cell must pad the row."""
    caminho = tmp_path / "d.parquet"
    pq.write_table(
        pa.table({"a": pa.array([1, 2]), "z": pa.array([None, None], type=pa.string())}),
        caminho,
    )
    destino = tmp_path / "r.xlsx"
    gerar.gerar_xlsx(
        str(caminho),
        {
            "nome": "L",
            "fonte": {"pipeline": str(uuid.uuid4()), "no": "n"},
            "aba": "D",
            "colunas": [{"coluna": "a"}, {"coluna": "z"}],
            "carimbo": False,
        },
        destino,
        carimbo=[],
    )
    ws = load_workbook(destino, read_only=True)["D"]
    linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    assert linhas[0] == ["a", "z"]
    assert linhas[1] == [1]  # the trailing empty cell was not written
    assert linhas[1] + [None] * (len(linhas[0]) - len(linhas[1])) == [1, None]
