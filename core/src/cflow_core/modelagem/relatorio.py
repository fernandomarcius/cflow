"""Stamped report: materializes a page's query into an IMMUTABLE record, with the trace chain
(report -> page/version -> views -> cubes -> execution -> nodes) via `trace`."""

from __future__ import annotations

import io
import json
import re
import uuid

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncConnection

from ..audit import Acao, registrar
from ..errors import ErroAplicacao
from ..timezone import para_exibicao
from . import consulta


async def gerar(
    conn: AsyncConnection,
    pagina_id: uuid.UUID,
    *,
    filtros: list[dict] | None,
    gerado_por: uuid.UUID,
    ip=None,
    request_id=None,
    execucao_id=None,
) -> dict:
    res = await consulta.consultar_pagina(
        conn, pagina_id, filtros_extra=filtros, execucao_id=execucao_id
    )
    trace = [
        {"view_id": b["view_id"], "origem": b["resultado"].get("origem")} for b in res["blocos"]
    ]
    rid = uuid.uuid4()
    await conn.execute(
        text(
            "INSERT INTO relatorios (id, pagina_id, pagina_versao, gerado_por, filtros, conteudo, trace) "
            "VALUES (:id, :p, :v, :g, CAST(:f AS jsonb), CAST(:c AS jsonb), CAST(:t AS jsonb))"
        ),
        {
            "id": rid,
            "p": pagina_id,
            "v": res["versao"],
            "g": gerado_por,
            "f": json.dumps(filtros or [], default=str),
            "c": json.dumps(res["blocos"], default=str),
            "t": json.dumps(trace, default=str),
        },
    )
    await registrar(
        conn,
        acao=Acao.EXPORTAR,
        recurso_tipo="relatorio",
        recurso_id=rid,
        ator_id=gerado_por,
        antes=None,
        depois={"pagina_id": str(pagina_id), "versao": res["versao"]},
        request_id=request_id,
        ip=ip,
    )
    return {
        "id": str(rid),
        "pagina": res["pagina"],
        "versao": res["versao"],
        "blocos": res["blocos"],
        "trace": trace,
    }


async def obter(conn: AsyncConnection, relatorio_id: uuid.UUID) -> dict:
    r = (
        (
            await conn.execute(
                text(
                    "SELECT r.id, r.pagina_id, r.pagina_versao, r.gerado_por, r.gerado_em, r.filtros, "
                    "r.conteudo, r.trace, p.nome AS pagina, u.nome AS gerado_por_nome "
                    "FROM relatorios r JOIN paginas p ON p.id = r.pagina_id "
                    "LEFT JOIN usuarios u ON u.id = r.gerado_por WHERE r.id = :id"
                ),
                {"id": relatorio_id},
            )
        )
        .mappings()
        .first()
    )
    if r is None:
        raise ErroAplicacao(codigo="nao_encontrado", mensagem="Report not found.", status=404)
    return {
        "id": str(r["id"]),
        "pagina_id": str(r["pagina_id"]),
        "pagina": r["pagina"],
        "versao": r["pagina_versao"],
        "gerado_em": para_exibicao(r["gerado_em"]).isoformat() if r["gerado_em"] else None,
        "criado_por": r["gerado_por_nome"],
        "filtros": r["filtros"],
        "blocos": r["conteudo"],
        "trace": r["trace"],
    }


async def listar(conn: AsyncConnection) -> list[dict]:
    """List the stamped reports (most recent first). `criado_em` in the display time zone."""
    rows = (
        (
            await conn.execute(
                text(
                    "SELECT r.id, p.nome AS pagina, r.pagina_versao, r.gerado_em, u.nome AS criado_por "
                    "FROM relatorios r JOIN paginas p ON p.id = r.pagina_id "
                    "LEFT JOIN usuarios u ON u.id = r.gerado_por ORDER BY r.gerado_em DESC"
                )
            )
        )
        .mappings()
        .all()
    )
    return [
        {
            "id": str(r["id"]),
            "pagina": r["pagina"],
            "versao": r["pagina_versao"],
            "criado_em": para_exibicao(r["gerado_em"]).isoformat() if r["gerado_em"] else None,
            "criado_por": r["criado_por"],
        }
        for r in rows
    ]


def _nome_aba(titulo: str | None, indice: int) -> str:
    """Sanitize the block title into an Excel sheet name (<=31 chars, no forbidden characters)."""
    bruto = (titulo or f"Block {indice + 1}").strip() or f"Block {indice + 1}"
    limpo = re.sub(r"[\[\]:*?/\\]", " ", bruto).strip()
    return (limpo or f"Block {indice + 1}")[:31]


def blocos_para_xlsx(blocos: list[dict]) -> bytes:
    """Build an .xlsx from the frozen blocks: one sheet per block, header = columns,
    rows = rows. Uses openpyxl (permissive license, works air-gapped). No data -> one empty sheet."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    vistos: dict[str, int] = {}
    for i, b in enumerate(blocos or []):
        res = b.get("resultado", {}) if isinstance(b, dict) else {}
        cols = list(res.get("colunas", []) or [])
        nome = _nome_aba(b.get("titulo") if isinstance(b, dict) else None, i)
        # De-duplicate sheet names (Excel requires uniqueness): "Name", "Name (2)"...
        if nome in vistos:
            vistos[nome] += 1
            sufixo = f" ({vistos[nome]})"
            nome = nome[: 31 - len(sufixo)] + sufixo
        else:
            vistos[nome] = 1
        ws = wb.create_sheet(title=nome)
        if cols:
            ws.append([str(c) for c in cols])
            for linha in res.get("linhas", []) or []:
                ws.append([linha.get(c) if isinstance(linha, dict) else None for c in cols])
    if not wb.sheetnames:
        wb.create_sheet(title="Report")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
