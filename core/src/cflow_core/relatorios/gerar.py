"""Node Parquet snapshot -> .xlsx file, via the layout. openpyxl in `write_only` (the repository's
only xlsx library), reading the Parquet in batches -- a million-row table does not fit in memory
as a list of dicts.

What goes into the cell is the value in the Parquet. NULL becomes an empty cell (never `0`), a
number stays a number (never `/100`), a date stays a date. The layout's `formato` is the cell's
`number_format`: it changes how Excel SHOWS the value, not what it stores -- this is the layer's
boundary inside this module.

Above the Excel cap (1,048,576 rows per sheet, header included) the data continues on the next
sheet -- `<sheet> (parte 2)`, same column order -- instead of being silently truncated. The
`Carimbo` sheet answers "where did this file come from?" for whoever only has the spreadsheet.
"""

from __future__ import annotations

import json
import math
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from functools import partial
from pathlib import Path

from openpyxl.styles.numbers import is_date_format

from ..errors import ErroAplicacao
from ..timezone import FUSO_SAO_PAULO
from . import celulas_xlsx
from . import layout as _layout

LOTE_LINHAS = 10_000
# What ClosedXML writes when it rewrites a copied sheet: the font completed with its defaults and
# the default column width (`<sheetFormatPr defaultColWidth="9.140625" defaultRowHeight="15"/>`,
# `<col min="1" max="16384" width="9.140625" customWidth="1"/>`, header Calibri 11 bold
# `rgb=FF000000`, family 2).
_FONTE_CLOSEDXML = "Calibri"
_TAMANHO_CLOSEDXML = 11
_COR_CLOSEDXML = "FF000000"
_FAMILIA_CLOSEDXML = 2
_LARGURA_CLOSEDXML = 9.140625
_ALTURA_CLOSEDXML = 15
_ULTIMA_COLUNA_EXCEL = 16384
ABA_CARIMBO = _layout.ABA_CARIMBO
# DATA rows per sheet: the Excel cap minus the header row.
TETO_LINHAS_DADO = _layout.TETO_LINHAS_EXCEL - 1


def _celula(valor: object) -> object:
    """Parquet value -> value openpyxl writes. Types Excel has (number, text, date, boolean) pass
    through. The rest is what Excel does NOT have, and the choice is stated here: NaN/inf do not
    exist in a cell (they become empty, like NULL); a timestamp with a time zone is converted to the
    reporting time zone and loses the offset, because Excel stores no zone; bytes come out as hex;
    lists/objects come out as JSON. None of these changes a number."""
    if valor is None:
        return None
    if isinstance(valor, bool | int | str | Decimal):
        return valor
    if isinstance(valor, float):
        return None if (math.isnan(valor) or math.isinf(valor)) else valor
    if isinstance(valor, datetime):
        if valor.tzinfo is not None:
            return valor.astimezone(FUSO_SAO_PAULO).replace(tzinfo=None)
        return valor
    if isinstance(valor, date):
        return valor
    if isinstance(valor, bytes | bytearray | memoryview):
        return bytes(valor).hex()
    if isinstance(valor, list | dict | tuple):
        return json.dumps(valor, ensure_ascii=False, default=str)
    return str(valor)


def _valor_closedxml(v: object, data: bool) -> object:
    """The value ClosedXML rewrites, on openpyxl's cell-by-cell path: empty text becomes a blank
    cell and the number goes through the SAME 15-digit rounding
    (`celulas_xlsx.numero_closedxml`). Here openpyxl still serializes the double with `%.16g` -- the
    double is the rewritten file's, the text may not be; the exact text is the columnar writer's."""
    if v == "" and type(v) is str:
        return None
    if type(v) in (int, float, Decimal):
        return float(celulas_xlsx.numero_closedxml(v, data))
    return v


def _larguras_closedxml(ws, larguras: list) -> None:
    """`<col>` as ClosedXML declares it: every column up to Excel's last with its default width
    (9.140625, `customWidth`), and the layout's width where there is one."""
    from openpyxl.worksheet.dimensions import ColumnDimension, SheetFormatProperties

    ws.sheet_format = SheetFormatProperties(
        baseColWidth=None, defaultColWidth=_LARGURA_CLOSEDXML, defaultRowHeight=_ALTURA_CLOSEDXML
    )
    por_coluna = [float(w) if w else _LARGURA_CLOSEDXML for w in larguras]
    faixas: list[list] = []  # [start, end, width]
    for i, w in enumerate(por_coluna, start=1):
        if faixas and faixas[-1][2] == w and faixas[-1][1] == i - 1:
            faixas[-1][1] = i
        else:
            faixas.append([i, i, w])
    if faixas and faixas[-1][2] == _LARGURA_CLOSEDXML:
        faixas[-1][1] = _ULTIMA_COLUNA_EXCEL
    else:
        faixas.append([len(por_coluna) + 1, _ULTIMA_COLUNA_EXCEL, _LARGURA_CLOSEDXML])
    for inicio, fim, w in faixas:
        letra = _letra_coluna(inicio)
        ws.column_dimensions[letra] = ColumnDimension(ws, index=letra, width=w, min=inicio, max=fim)


def _nome_parte(aba: str, parte: int, total: int) -> str:
    if total <= 1:
        return aba
    sufixo = f" (parte {parte})"
    return aba[: 31 - len(sufixo)] + sufixo


def colunas_do_parquet(caminho: str | Path) -> list[str]:
    import pyarrow.parquet as pq

    return list(pq.ParquetFile(str(caminho)).schema_arrow.names)


def linhas_do_parquet(caminho: str | Path) -> int:
    import pyarrow.parquet as pq

    return int(pq.ParquetFile(str(caminho)).metadata.num_rows)


def _celula_formatada(
    ws, valor, *, formato=None, negrito=False, centro=False, tamanho=None, closedxml=False
):
    """Styled cell. In `write_only` the style goes on the cell, not on the sheet afterwards.

    `closedxml`: the font comes out COMPLETE, as ClosedXML rewrites it when copying the sheet
    (Calibri, explicit `FF000000` color, family 2): a lone `<b/>` font becomes this. See
    `gerar_xlsx_agrupado`."""
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font

    cel = WriteOnlyCell(ws, value=valor)
    if formato:
        cel.number_format = formato
    if closedxml and (negrito or tamanho):
        cel.font = Font(
            name=_FONTE_CLOSEDXML,
            sz=tamanho or _TAMANHO_CLOSEDXML,
            b=bool(negrito),
            color=_COR_CLOSEDXML,
            family=_FAMILIA_CLOSEDXML,
        )
    elif negrito or tamanho:
        cel.font = Font(bold=bool(negrito), size=tamanho) if tamanho else Font(bold=True)
    if centro:
        cel.alignment = Alignment(horizontal="center")
    return cel


_PLACEHOLDER = re.compile(r"\{\{(param|coluna)\.([A-Za-z_][A-Za-z0-9_]*)(?::([^}]+))?\}\}")


def valores_do_preambulo(arquivo, definicao: dict) -> dict:
    """The value of each column the preamble cites via `{{coluna.x}}`, read from the snapshot ITSELF.

    The column must be CONSTANT: the preamble has one cell, and picking a row among different
    values would be the layout deciding a number. More than one value is an error naming the
    column; a NULL-only column resolves empty, like a parameter placeholder with no value."""
    nomes = _layout.colunas_do_preambulo(definicao)
    if not nomes:
        return {}
    import pyarrow.compute as pc

    tabela = arquivo.read(columns=nomes)
    valores: dict = {}
    for nome in nomes:
        distintos = pc.unique(tabela.column(nome).drop_null()).to_pylist()
        if len(distintos) > 1:
            raise ErroAplicacao(
                codigo="layout_valor_nao_constante",
                mensagem=(
                    f"Column '{nome}', cited in the layout preamble, has {len(distintos)} "
                    "different values in the snapshot. The preamble shows ONE value, and the "
                    "layout does not choose among them: a pipeline node decides the number."
                ),
                status=400,
            )
        valores[nome] = distintos[0] if distintos else None
    return valores


def _resolver(texto: object, parametros: dict, colunas: dict | None = None) -> object:
    """Replaces `{{param.nome}}` with the run's value -- the SAME convention as the e-mail node --
    and `{{coluna.nome}}` with that column's value in the snapshot (`valores_do_preambulo`).

    What does not resolve becomes empty, as there: a mistyped placeholder leaking raw into the sheet
    title is worse than a gap, because it looks like a real name. A `segredo` parameter never
    reaches here (the persisted value is `***`), and the layout only reaches what it declares.
    """
    if not isinstance(texto, str) or "{{" not in texto:
        return texto

    def _um(m: re.Match) -> str:
        origem = parametros if m.group(1) == "param" else (colunas or {})
        valor = origem.get(m.group(2))
        if valor in (None, ""):
            return ""
        formato = m.group(3)
        if formato:
            # `{{param.period_start:%m_%Y}}` -- a display FORMAT for a date, which is what this
            # layer does per column all the time. Not a calculation: the value does not change,
            # only how it is written. A date that does not parse comes out as given, never empty.
            try:
                return datetime.fromisoformat(str(valor)).strftime(formato)
            except ValueError:
                return str(valor)
        return str(valor)

    return _PLACEHOLDER.sub(_um, texto)


def _valor_fixo(f: dict, parametros: dict, colunas: dict | None = None) -> object:
    """Value of a fixed row, already resolved. `escala` is PRESENTATION, not a rule: an average
    arrives in POINTS (47.41) and the cell is formatted as a percentage, which Excel reads as a
    fraction -- without the scale the file would show 4741%."""
    bruto = _resolver(f.get("texto", f.get("valor")), parametros, colunas)
    escala = f.get("escala")
    if escala is None or bruto in (None, ""):
        return bruto
    try:
        return float(bruto) * float(escala)
    except (TypeError, ValueError):
        return bruto


def _preambulo(
    ws, definicao, rotulos, colunas, primeira, ultima, largura, tamanho, parametros, valores=None
):
    """The rows BEFORE the data (the named fixed-block exception).

    In `write_only` there is no going back: each row is written once, in order. So the preamble is
    built by COORDINATE -- title, loose blocks, header and the totals row, which sits ABOVE the
    data and references a range that only exists afterwards.
    """
    titulo = definicao.get("titulo") or {}
    fixas: dict[int, list[dict]] = {}
    for f in definicao.get("linhas_fixas") or []:
        fixas.setdefault(int(f["linha"]), []).append(f)
    totais = definicao.get("totais") or {}
    linha_cab = int(definicao.get("linha_cabecalho") or (primeira - 1))
    formatos = _layout.formatos_do_layout(definicao)

    saida = []
    for numero in range(1, primeira):
        linha = [None] * largura
        if numero == 1 and titulo.get("texto"):
            linha[0] = _celula_formatada(
                ws,
                _resolver(titulo["texto"], parametros, valores),
                negrito=titulo.get("negrito", True),
                centro=True,
                tamanho=tamanho,
            )
        for f in fixas.get(numero, []):
            col = int(f.get("coluna", 1)) - 1
            if 0 <= col < largura:
                linha[col] = _celula_formatada(
                    ws,
                    _valor_fixo(f, parametros, valores),
                    formato=_layout.FORMATOS.get(f.get("formato") or ""),
                    negrito=f.get("negrito", False),
                    centro=f.get("centralizado", False),
                    tamanho=tamanho,
                )
        if numero == linha_cab:
            for i, rot in enumerate(rotulos):
                if rot and i < largura:
                    linha[i] = _celula_formatada(ws, rot, negrito=True, tamanho=tamanho)
        if totais and numero == int(totais.get("linha", -1)):
            for nome in totais.get("colunas") or []:
                i = colunas.index(nome)
                letra = _letra_coluna(i + 1)
                linha[i] = _celula_formatada(
                    ws,
                    f"=SUM({letra}{primeira}:{letra}{ultima})",
                    formato=formatos[i],
                    negrito=totais.get("negrito", True),
                    tamanho=tamanho,
                )
        saida.append(linha)
    return saida


@dataclass(frozen=True)
class ParteXlsx:
    """One SOURCE of the workbook: the snapshot, the single-source layout that organizes it and, in
    a grouped file, the sheet name the grouped layout gives it (`None` = the layout's own `aba`)."""

    caminho_parquet: str | Path
    definicao: dict
    aba: str | None = None
    parametros: dict | None = None


class _Pasta:
    """What belongs to the WHOLE workbook, not one sheet: the `write_only` workbook, the styles (a
    single registry, in the order cells are born -- it numbers `styles.xml`), the sheet names
    already used and each sheet's data rows (injected into the zip after `save`)."""

    def __init__(self, escritor: str, *, closedxml: bool = False) -> None:
        from openpyxl import Workbook

        self.wb = Workbook(write_only=True)
        self.colunar = escritor == "colunar"
        self.closedxml = closedxml
        self.estilos = celulas_xlsx.Estilos()
        self.abas: list[str] = []
        self.por_aba: list[tuple] = []  # (sheet, its rows)
        # (sheet, `<dimension ref>`): ClosedXML declares the used range; openpyxl in `write_only`
        # does not. The sheet's path only exists after `save`.
        self.dimensoes: list[tuple] = []

    def criar_aba(self, nome: str):
        # Excel compares sheet names ignoring case, and openpyxl would silently rename the repeated
        # one (`SALES1`): a sheet under another name is the wrong file looking right.
        if any(nome.casefold() == a.casefold() for a in self.abas):
            raise ErroAplicacao(
                codigo="layout_aba_repetida",
                mensagem=f"Two sheets of the workbook would have the same name '{nome}'. Excel does not "
                "distinguish case in sheet names; rename one of them in the layout.",
                status=400,
            )
        ws = self.wb.create_sheet(title=nome)
        self.abas.append(nome)
        return ws

    def descartar(self, *, falhou: bool = False) -> None:
        for _, linhas in self.por_aba:
            linhas.fechar().unlink(missing_ok=True)
        if not falhou:
            return
        # Error midway (repeated sheet, illegal character): the `write_only` sheets already open
        # each hold a temporary file that only `save` would delete.
        for ws in self.wb.worksheets:
            try:
                if not ws.closed:
                    ws.close()
                saida = getattr(getattr(ws, "_writer", None), "out", None)
                if saida:
                    Path(saida).unlink(missing_ok=True)
            except Exception:  # noqa: BLE001, S110 -- error cleanup; the original error is the one that propagates
                pass


def _preparar(parte: ParteXlsx):
    """Checks the part BEFORE any sheet is written: layout, columns against the snapshot and the
    preamble's constant value. In a grouped file, a third part with a missing column fails before
    the first becomes a sheet -- no half-written file."""
    import pyarrow.parquet as pq

    _layout.validar(parte.definicao)
    if _layout.eh_agrupado(parte.definicao):
        raise ErroAplicacao(
            codigo="layout_agrupado_aninhado",
            mensagem="A sheet of the grouped layout points to another grouped layout; each sheet "
            "organizes the snapshot of ONE node.",
            status=400,
        )
    arquivo = pq.ParquetFile(str(parte.caminho_parquet))
    _layout.validar_contra_snapshot(parte.definicao, list(arquivo.schema_arrow.names))
    return arquivo, valores_do_preambulo(arquivo, parte.definicao)


def _escrever_fonte(
    pasta: _Pasta,
    arquivo,
    valores_preambulo: dict,
    parte: ParteXlsx,
    *,
    teto_linhas: int,
    progresso: Callable[[int], None] | None,
    ja_escritas: int,
) -> int:
    """The sheets of ONE source: header, data in batches and, above the cap, `<sheet> (parte k)`.
    The output is the same with one source or several -- which makes the grouped file equal, sheet
    by sheet, to each source's own file. Returns the rows written.

    With `pasta.closedxml`, the way of writing is ClosedXML's when rewriting a sheet (see
    `gerar_xlsx_agrupado`); each cell's value does not change."""
    from openpyxl.cell import WriteOnlyCell

    definicao = parte.definicao
    parametros = parte.parametros
    colunas = _layout.colunas_do_layout(definicao)
    rotulos = _layout.rotulos_do_layout(definicao)
    rotulos2 = _layout.rotulos2_do_layout(definicao)
    linhas_cabecalho = int(definicao.get("linhas_cabecalho", 1))
    # Nameless, dataless trailing columns some files must keep. The only way for openpyxl in
    # `write_only` to create the column is to write an empty string: a `None` cell at the end of
    # a row is omitted, and the column would not exist in the file.
    vazias = [""] * int(definicao.get("colunas_vazias_no_fim", 0))
    # Fixed block: the data starts on a row > 1, and what comes before (merged title, loose
    # values, header and totals) is built by coordinate. See `_preambulo`.
    primeira_dado = int(definicao.get("primeira_linha_dado") or 0)
    tamanho_fonte = definicao.get("fonte_tamanho")
    negrito_por_coluna = [bool(c.get("negrito")) for c in definicao.get("colunas", [])]
    formatos = _layout.formatos_do_layout(definicao)
    larguras = [c.get("largura") for c in definicao.get("colunas", [])]
    total = int(arquivo.metadata.num_rows)
    partes = max(1, math.ceil(total / teto_linhas)) if total else 1
    aba_base = _layout.nome_aba(parte.aba if parte.aba is not None else definicao.get("aba"))
    cabecalho_fixo = definicao.get("cabecalho_fixo", True)
    autofiltro = definicao.get("autofiltro", True)
    ultima_coluna = _letra_coluna(len(colunas) + len(vazias))

    ws = None
    numero = 0
    na_aba = 0  # data rows already written on the current sheet
    colunar = pasta.colunar
    closedxml = pasta.closedxml
    formatada = partial(_celula_formatada, closedxml=closedxml)
    datas = [bool(f) and is_date_format(f) for f in formatos]
    linhas_aba: celulas_xlsx.LinhasDaAba | None = None

    def _abrir_aba():
        nonlocal ws, numero, na_aba, linhas_aba
        numero += 1
        ws = pasta.criar_aba(_nome_parte(aba_base, numero, partes))
        # Width and frozen pane before the first row: in `write_only` anything set afterwards does
        # not reach the file.
        if closedxml:
            _larguras_closedxml(ws, larguras)
        else:
            for i, largura in enumerate(larguras, start=1):
                if largura:
                    ws.column_dimensions[_letra_coluna(i)].width = float(largura)
        if titulo_mesclar := (definicao.get("titulo") or {}).get("ate_coluna"):
            ws.merged_cells.ranges.add(f"A1:{titulo_mesclar}1")
        if cabecalho_fixo:
            ws.freeze_panes = "A2"
        if primeira_dado > 1:
            largura = len(colunas) + len(vazias)
            ultima = primeira_dado + min(total, teto_linhas) - 1
            for linha_pre in _preambulo(
                ws,
                definicao,
                rotulos,
                colunas,
                primeira_dado,
                ultima,
                largura,
                tamanho_fonte,
                parametros or {},
                valores_preambulo,
            ):
                ws.append(linha_pre)
        else:
            # BOLD header, on both rows when there are two.
            def _cab(textos):
                # The trailing empty column is bold too: the header style spans the whole range.
                # In ClosedXML it is a BLANK cell with the style (not empty text).
                return [
                    formatada(
                        ws,
                        None if closedxml and t == "" else t,
                        negrito=True,
                        tamanho=tamanho_fonte,
                    )
                    for t in [*textos, *([None] * len(vazias) if closedxml else vazias)]
                ]

            ws.append(_cab(rotulos))
            if linhas_cabecalho == 2:
                ws.append(_cab(rotulos2))
        na_aba = 0
        if colunar:
            linhas_aba = celulas_xlsx.LinhasDaAba(
                ws,
                list(zip(formatos, negrito_por_coluna, strict=True)),
                tamanho=tamanho_fonte,
                vazias=len(vazias),
                primeira_linha=primeira_dado if primeira_dado > 1 else linhas_cabecalho + 1,
                estilos=pasta.estilos,
                celula=_celula,
                celula_formatada=formatada,
                closedxml=closedxml,
            )
            pasta.por_aba.append((ws, linhas_aba))

    def _fechar_aba():
        if ws is not None and autofiltro:
            ws.auto_filter.ref = f"A1:{ultima_coluna}{na_aba + linhas_cabecalho}"
        if ws is not None and closedxml:
            # The used range: the header (or preamble) spans the full width, including the
            # trailing empty column, and the data runs to the last row written.
            ultima = (primeira_dado - 1 if primeira_dado > 1 else linhas_cabecalho) + na_aba
            pasta.dimensoes.append((ws, f"A1:{ultima_coluna}{ultima}"))

    _abrir_aba()
    escritas = 0
    for lote in arquivo.iter_batches(batch_size=LOTE_LINHAS, columns=colunas):
        series = [lote.column(i).to_pylist() for i in range(lote.num_columns)]
        if colunar:
            # The batch is sliced at the sheet cap: a new sheet only opens when there is a row for
            # it, as in the cell-by-cell loop (a total that is a multiple of the cap creates no
            # empty sheet).
            linhas = list(zip(*series, strict=True))
            i = 0
            while i < len(linhas):
                if na_aba >= teto_linhas:
                    _fechar_aba()
                    _abrir_aba()
                k = min(len(linhas) - i, teto_linhas - na_aba)
                linhas_aba.escrever(linhas[i : i + k])
                na_aba += k
                i += k
        else:
            for valores in zip(*series, strict=True):
                if na_aba >= teto_linhas:
                    _fechar_aba()
                    _abrir_aba()
                linha = []
                for valor, fmt, negrito, data in zip(
                    valores, formatos, negrito_por_coluna, datas, strict=True
                ):
                    v = _celula(valor)
                    if closedxml:
                        v = _valor_closedxml(v, data)
                    if fmt is None and not negrito and not tamanho_fonte:
                        linha.append(v)
                    elif negrito or tamanho_fonte:
                        linha.append(
                            formatada(ws, v, formato=fmt, negrito=negrito, tamanho=tamanho_fonte)
                        )
                    else:
                        # The format goes on the cell EVEN when empty: applying it by RANGE would
                        # leave a date column with gaps as `General` on the null rows only; the
                        # whole column must be formatted.
                        cel = WriteOnlyCell(ws, value=v)
                        cel.number_format = fmt
                        linha.append(cel)
                # ClosedXML: the trailing empty column has no cell in data rows.
                ws.append(linha if closedxml else [*linha, *vazias])
                na_aba += 1
        escritas += lote.num_rows
        if progresso is not None:
            progresso(ja_escritas + escritas)
    _fechar_aba()
    return escritas


def _gravar_pasta(pasta: _Pasta, destino: str | Path, carimbo, com_carimbo: bool) -> Path:
    if com_carimbo:
        wc = pasta.criar_aba(ABA_CARIMBO)
        wc.column_dimensions["A"].width = 28.0
        wc.column_dimensions["B"].width = 60.0
        wc.append(["campo", "valor"])
        for chave, valor in carimbo or []:
            wc.append([chave, _celula(valor)])

    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    pasta.wb.save(str(destino))
    dimensoes = {planilha.path.lstrip("/"): ref for planilha, ref in pasta.dimensoes}
    if pasta.colunar or dimensoes:
        # The data rows go in now, before each sheet's `</sheetData>` (and ClosedXML's
        # `<dimension>`, when requested).
        celulas_xlsx.injetar(
            destino,
            {planilha.path.lstrip("/"): linhas.fechar() for planilha, linhas in pasta.por_aba},
            dimensoes=dimensoes,
            # Same sheets that get the <dimension>: the grouped file's data sheets (ClosedXML mode).
            impressao=set(dimensoes) if pasta.closedxml else None,
        )
    return destino


def _escritor_efetivo(escritor: str) -> str:
    if escritor not in ("colunar", "openpyxl"):
        raise ValueError(f"unknown writer: {escritor}")
    if escritor == "colunar" and not celulas_xlsx.compativel():
        return "openpyxl"
    return escritor


def gerar_xlsx(
    caminho_parquet: str | Path,
    definicao: dict,
    destino: str | Path,
    *,
    carimbo: list[tuple[str, object]] | None = None,
    parametros: dict | None = None,
    teto_linhas: int = TETO_LINHAS_DADO,
    progresso: Callable[[int], None] | None = None,
    escritor: str = "colunar",
) -> dict:
    """Writes the .xlsx to `destino` and returns `{linhas, abas, bytes}`.

    `teto_linhas` is the maximum number of DATA rows per sheet (a parameter so tests can lower it --
    Excel's real cap is 1,048,575 data rows, and no test should generate a million rows).
    `progresso(linhas_escritas)` is called per batch, so the UI can show how much is done.

    `escritor`: `colunar` (default) writes the data rows through `celulas_xlsx`, without an
    openpyxl object per cell; `openpyxl` is the cell-by-cell path -- kept as the REFERENCE the
    tests use to prove both files come out identical, byte for byte.
    """
    escritor = _escritor_efetivo(escritor)
    if teto_linhas < 1:
        raise ValueError("teto_linhas must be >= 1")
    parte = ParteXlsx(caminho_parquet, definicao, None, parametros)
    arquivo, valores = _preparar(parte)
    pasta = _Pasta(escritor)
    try:
        escritas = _escrever_fonte(
            pasta,
            arquivo,
            valores,
            parte,
            teto_linhas=teto_linhas,
            progresso=progresso,
            ja_escritas=0,
        )
        destino = _gravar_pasta(pasta, destino, carimbo, bool(definicao.get("carimbo", True)))
    except BaseException:
        pasta.descartar(falhou=True)
        raise
    pasta.descartar()
    return {"linhas": escritas, "abas": list(pasta.abas), "bytes": destino.stat().st_size}


def gerar_xlsx_agrupado(
    partes: list[ParteXlsx],
    destino: str | Path,
    *,
    carimbo: list[tuple[str, object]] | None = None,
    com_carimbo: bool = False,
    teto_linhas: int = TETO_LINHAS_DADO,
    progresso: Callable[[int], None] | None = None,
    escritor: str = "colunar",
    closedxml: bool = True,
) -> dict:
    """ONE .xlsx with the sheets of SEVERAL sources, in the given order.

    Each part is written by the SAME `_escrever_fonte` as a single-source layout -- header, format,
    bold, empty columns, width, pane, filter, preamble and `<sheet> (parte k)` above the cap --
    only with the sheet name the grouped layout gives it. It re-reads nobody's xlsx: the data comes
    from the snapshot, as in each source's own file. `progresso` receives the running sum of rows
    written.

    `closedxml` (default): the file comes out in the SHAPE a ClosedXML rewrite produces, as when
    sheets are copied from each source file and rewritten. Only the way of writing changes, never
    the data: default column width and `sheetFormatPr`, `<dimension>`, a complete header font
    (Calibri 11, color `FF000000`), a blank cell where openpyxl would write empty text (the
    trailing empty column and a data `''`) and numbers with `double.ToString("G15")` text -- the
    15-digit rounding ClosedXML applies on rewrite (`celulas_xlsx.numero_closedxml`). `False` =
    each sheet identical, byte for byte, to the single source's own sheet.

    Returns `{linhas, abas, bytes, partes: [{aba, linhas, abas}]}`."""
    escritor = _escritor_efetivo(escritor)
    if teto_linhas < 1:
        raise ValueError("teto_linhas must be >= 1")
    if not partes:
        raise ValueError("a grouped file needs at least one part")
    preparadas = [(p, *_preparar(p)) for p in partes]
    pasta = _Pasta(escritor, closedxml=closedxml)
    try:
        total = 0
        por_parte = []
        for parte, arquivo, valores in preparadas:
            antes = len(pasta.abas)
            n = _escrever_fonte(
                pasta,
                arquivo,
                valores,
                parte,
                teto_linhas=teto_linhas,
                progresso=progresso,
                ja_escritas=total,
            )
            total += n
            por_parte.append({"aba": parte.aba, "linhas": n, "abas": pasta.abas[antes:]})
        destino = _gravar_pasta(pasta, destino, carimbo, com_carimbo)
    except BaseException:
        pasta.descartar(falhou=True)
        raise
    pasta.descartar()
    return {
        "linhas": total,
        "abas": list(pasta.abas),
        "bytes": destino.stat().st_size,
        "partes": por_parte,
    }


def _letra_coluna(n: int) -> str:
    """1 -> A, 27 -> AA (what openpyxl's `get_column_letter` does, without importing it)."""
    s = ""
    while n > 0:
        n, resto = divmod(n - 1, 26)
        s = chr(65 + resto) + s
    return s or "A"


def momento_carimbo(momento: datetime | None = None) -> str:
    """Stamp instant in the reporting time zone with an explicit offset (repository rule)."""
    m = momento or datetime.now(tz=UTC)
    if m.tzinfo is None:
        m = m.replace(tzinfo=UTC)
    return m.astimezone(FUSO_SAO_PAULO).isoformat(timespec="seconds")
