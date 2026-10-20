"""File source connector: CSV, Excel (xlsx) and Parquet. Reads from a path reachable by the worker.

config: {caminho, formato? ('csv'|'xlsx'|'parquet', otherwise inferred from the extension), aba?
(sheet name for xlsx, otherwise the active one)}.
Reading is synchronous; it runs in a thread (asyncio.to_thread) so the event loop is not blocked.
"""

from __future__ import annotations

import asyncio
import csv
import os

from .base import ResultadoTeste, serial


def _formato(config: dict) -> str:
    f = config.get("formato")
    if f:
        return f
    ext = os.path.splitext(config.get("caminho", ""))[1].lower().lstrip(".")
    return {"csv": "csv", "xlsx": "xlsx", "xls": "xlsx", "parquet": "parquet", "pq": "parquet"}.get(
        ext, ext
    )


def _ler(config: dict, limite: int | None) -> tuple[list[str], list[dict]]:
    """Returns `(colunas, linhas)`. The header is always known (DictReader.fieldnames / first xlsx
    row / parquet `cur.description`), so columns exist even with 0 data rows.
    limite None = read everything (execution)."""
    caminho = config["caminho"]
    if not os.path.isfile(caminho):
        raise FileNotFoundError(caminho)
    fmt = _formato(config)

    if fmt == "csv":
        with open(caminho, newline="", encoding="utf-8") as fp:
            leitor = csv.DictReader(fp)
            linhas = []
            for i, row in enumerate(leitor):
                if limite is not None and i >= limite:
                    break
                linhas.append({k: serial(v) for k, v in row.items()})
            colunas = list(leitor.fieldnames or [])  # CSV header (exists without data rows)
            return colunas, linhas

    if fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(caminho, read_only=True, data_only=True)
        aba = config.get("aba")  # specific sheet (structured Excel); otherwise the active one
        if aba and aba not in wb.sheetnames:
            wb.close()
            raise ValueError(f"sheet '{aba}' not found (sheets: {wb.sheetnames})")
        ws = wb[aba] if aba else wb.active
        it = ws.iter_rows(values_only=True)
        cabec = [str(c) for c in next(it, [])]  # first row = header (exists without data rows)
        linhas = []
        for i, row in enumerate(it):
            if limite is not None and i >= limite:  # None = read everything (execution)
                break
            linhas.append({c: serial(v) for c, v in zip(cabec, row, strict=False)})
        wb.close()
        return cabec, linhas

    if fmt == "parquet":
        import duckdb

        con = duckdb.connect()
        try:
            if limite is None:
                cur = con.execute("SELECT * FROM read_parquet(?)", [caminho])
            else:
                cur = con.execute("SELECT * FROM read_parquet(?) LIMIT ?", [caminho, int(limite)])
            cols = [d[0] for d in cur.description]  # parquet schema (exists without data rows)
            return cols, [
                {c: serial(v) for c, v in zip(cols, row, strict=False)} for row in cur.fetchall()
            ]
        finally:
            con.close()

    raise ValueError(f"unsupported format: {fmt}")


def _causa(exc: Exception) -> str:
    if isinstance(exc, FileNotFoundError):
        return "file not found"
    if isinstance(exc, PermissionError):
        return "no permission to read the file"
    if isinstance(exc, ValueError):
        return str(exc)
    return "failed to read the file"


class ConectorArquivo:
    async def testar(self, config: dict, senha: str | None) -> ResultadoTeste:
        try:
            await asyncio.to_thread(_ler, config, 1)
            return ResultadoTeste(ok=True)
        except Exception as exc:  # noqa: BLE001 — sanitised cause
            return ResultadoTeste(ok=False, causa=_causa(exc))

    async def schemas_tabelas(self, config: dict, senha: str | None) -> list[dict]:
        from ...errors import ErroAplicacao

        raise ErroAplicacao(
            codigo="nao_aplicavel",
            mensagem="File source has no schema browsing.",
            status=400,
        )

    async def colunas(
        self, config: dict, senha: str | None, schema: str, tabela: str
    ) -> list[dict]:
        colunas, linhas = await asyncio.to_thread(_ler, config, 1)
        if linhas:  # with data: type inferred from the first value
            return [{"nome": k, "tipo": type(v).__name__} for k, v in linhas[0].items()]
        return [
            {"nome": c, "tipo": "str"} for c in colunas
        ]  # without data: still lists the columns

    async def amostra(
        self, config: dict, senha: str | None, sql: str, limite: int, params: dict | None = None
    ) -> tuple[list[str], list[dict]]:
        # A file source does not run user SQL (it reads the file), so params do not apply.
        return await asyncio.to_thread(_ler, config, limite)
